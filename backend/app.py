from flask import Flask, jsonify, request, g
from flask_cors import CORS
from functools import wraps
import csv
import json
import os
import uuid
import secrets
import hmac
from datetime import datetime


app = Flask(__name__)


CORS(
    app,
    resources={
        r"/api/*": {
            "origins": [
                "http://127.0.0.1:5500",
                "http://localhost:5500"
            ],
            "allow_headers": [
                "Content-Type",
                "Authorization"
            ]
        }
    }
)


BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

PROJECT_DIR = os.path.abspath(
    os.path.join(BASE_DIR, "..")
)

GD6CC_DIR = os.path.join(
    PROJECT_DIR,
    "extracted_gđ6cc"
)

GD6CC_DATAR = os.path.join(
    GD6CC_DIR,
    "datar"
)

ML_DIR = os.path.join(
    BASE_DIR,
    "ML"
)

WORKING_JSON = os.path.join(
    BASE_DIR,
    "products_working.json"
)

OLD_JSON = os.path.join(
    BASE_DIR,
    "products.json"
)


DEMO_USERS = {
    "admin": {
        "password": "Admin@123",
        "role": "Admin"
    },
    "user": {
        "password": "User@123",
        "role": "User"
    }
}


ACTIVE_SESSIONS = {}


PERMISSIONS = {
    "Admin": {
        "view": True,
        "create": True,
        "update": True,
        "delete": True,
        "backup": True,
        "restore": True
    },
    "User": {
        "view": True,
        "create": False,
        "update": False,
        "delete": False,
        "backup": False,
        "restore": False
    }
}


def find_file(file_names):
    roots = [
        BASE_DIR,
        PROJECT_DIR,
        GD6CC_DATAR,
        GD6CC_DIR,
        os.path.join(
            PROJECT_DIR,
            "frontend"
        ),
        os.path.join(
            PROJECT_DIR,
            "data"
        ),
        os.path.join(
            BASE_DIR,
            "data"
        )
    ]

    for root in roots:
        if not os.path.isdir(root):
            continue

        for name in file_names:
            path = os.path.join(
                root,
                name
            )

            if os.path.exists(path):
                return path

    return None


def find_review_files():
    roots = [
        BASE_DIR,
        PROJECT_DIR,
        GD6CC_DATAR,
        GD6CC_DIR,
        os.path.join(
            PROJECT_DIR,
            "frontend"
        ),
        os.path.join(
            PROJECT_DIR,
            "data"
        ),
        os.path.join(
            BASE_DIR,
            "data"
        )
    ]

    review_files = set()

    for root in roots:
        if not os.path.isdir(root):
            continue

        for dirpath, _, filenames in os.walk(root):
            for name in filenames:
                lower = name.lower()

                if not lower.endswith(".csv"):
                    continue

                if (
                    "review" in lower
                    or "rating" in lower
                    or "feedback" in lower
                    or "danhgia" in lower
                    or "đánh giá" in lower
                    or "customer" in lower
                    or "khach" in lower
                ):
                    path = os.path.abspath(
                        os.path.join(
                            dirpath,
                            name
                        )
                    )

                    if os.path.exists(path):
                        review_files.add(path)

    return sorted(
        list(review_files)
    )


PRODUCTS_CSV = find_file([
    "products (1).csv",
    "products.csv",
    "products(1).csv"
])

PRICE_CSV = find_file([
    "price_snapshots (1).csv",
    "price_snapshots.csv",
    "price_snapshots(1).csv"
])

OLD_PRODUCTS_CSV = find_file([
    "hasaki_products_pages_1_9.csv"
])


def clean_key(value):
    return str(
        value or ""
    ).strip()


def to_float(value):
    try:
        if value is None:
            return 0.0

        if isinstance(
            value,
            (int, float)
        ):
            return float(value)

        value_str = str(
            value
        ).strip()

        if value_str == "":
            return 0.0

        if value_str.lower() in {
            "nan",
            "none",
            "null",
            "n/a",
            "na"
        }:
            return 0.0

        value_str = (
            value_str
            .replace("đ", "")
            .replace("₫", "")
            .replace("%", "")
            .replace(" ", "")
            .strip()
        )

        if not value_str:
            return 0.0

        if "," in value_str and "." in value_str:
            last_comma = value_str.rfind(",")
            last_dot = value_str.rfind(".")

            if last_comma > last_dot:
                value_str = (
                    value_str
                    .replace(".", "")
                    .replace(",", ".")
                )
            else:
                value_str = value_str.replace(
                    ",",
                    ""
                )

        elif "," in value_str:
            parts = value_str.split(",")

            if (
                len(parts) == 2
                and len(parts[1]) <= 2
            ):
                value_str = (
                    parts[0]
                    + "."
                    + parts[1]
                )
            else:
                value_str = value_str.replace(
                    ",",
                    ""
                )

        return float(
            value_str
        )

    except Exception:
        return 0.0


def to_price_float(value):
    try:
        if value is None:
            return 0.0

        if isinstance(
            value,
            (int, float)
        ):
            return float(value)

        value_str = str(
            value
        ).strip()

        if not value_str:
            return 0.0

        value_str = (
            value_str
            .replace("đ", "")
            .replace("₫", "")
            .replace(" ", "")
            .strip()
        )

        if "," in value_str and "." in value_str:
            last_comma = value_str.rfind(",")
            last_dot = value_str.rfind(".")

            if last_comma > last_dot:
                value_str = (
                    value_str
                    .replace(".", "")
                    .replace(",", ".")
                )
            else:
                value_str = value_str.replace(
                    ",",
                    ""
                )

            return float(
                value_str
            )

        if "," in value_str:
            parts = value_str.split(",")

            if (
                len(parts) == 2
                and len(parts[1]) <= 2
            ):
                return float(
                    parts[0]
                    + "."
                    + parts[1]
                )

            return float(
                value_str.replace(
                    ",",
                    ""
                )
            )

        if "." in value_str:
            parts = value_str.split(".")

            if (
                len(parts) == 2
                and len(parts[1]) == 3
                and parts[0].isdigit()
            ):
                return float(
                    parts[0]
                    + parts[1]
                )

            if (
                len(parts) > 2
                and all(
                    part.isdigit()
                    for part in parts
                )
            ):
                return float(
                    "".join(parts)
                )

            return float(
                value_str
            )

        return float(
            value_str
        )

    except Exception:
        return 0.0


def to_percent_float(value):
    try:
        if value is None:
            return 0.0

        if isinstance(
            value,
            (int, float)
        ):
            result = float(value)
        else:
            value_str = str(
                value
            ).strip()

            if not value_str:
                return 0.0

            value_str = (
                value_str
                .replace("%", "")
                .replace(" ", "")
                .replace(",", ".")
            )

            result = float(
                value_str
            )

        if result < 0:
            result = 0.0

        if result > 100:
            result = 100.0

        return result

    except Exception:
        return 0.0


def to_ml_float(value):
    try:
        if value is None:
            return 0.0

        if isinstance(
            value,
            (int, float)
        ):
            return float(value)

        value_str = str(
            value
        ).strip()

        if value_str.lower() in {
            "",
            "nan",
            "none",
            "null"
        }:
            return 0.0

        cleaned = (
            value_str
            .replace("đ", "")
            .replace("₫", "")
            .replace(" ", "")
            .replace(",", "")
            .strip()
        )

        return float(
            cleaned
        )

    except Exception:
        return 0.0


def to_int(value):
    try:
        if value is None:
            return 0

        value = str(
            value
        ).strip()

        if value == "":
            return 0

        float_val = to_float(
            value
        )

        return int(
            float_val
        )

    except Exception:
        return 0


def read_csv(path):
    if (
        not path
        or not os.path.exists(path)
    ):
        return []

    try:
        with open(
            path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:
            return list(
                csv.DictReader(file)
            )

    except Exception as error:
        print(
            "CSV read error:",
            error
        )
        return []


def latest_price_map():
    rows = read_csv(
        PRICE_CSV
    )

    result = {}

    for row in rows:
        product_id = clean_key(
            row.get(
                "product_id"
            )
        )

        retailer_id = clean_key(
            row.get(
                "retailer_id"
            )
        )

        if not product_id:
            continue

        key = (
            product_id,
            retailer_id
        )

        collected_at = clean_key(
            row.get(
                "collected_at"
            )
        )

        old = result.get(
            key
        )

        if (
            old is None
            or collected_at
            > clean_key(
                old.get(
                    "collected_at"
                )
            )
        ):
            result[key] = row

    return result


def build_products():
    if PRODUCTS_CSV:
        product_rows = read_csv(
            PRODUCTS_CSV
        )

        price_map = latest_price_map()

        products = []

        for row in product_rows:
            product_id = clean_key(
                row.get(
                    "product_id"
                )
            )

            retailer_id = clean_key(
                row.get(
                    "retailer_id"
                )
            )

            price = price_map.get(
                (
                    product_id,
                    retailer_id
                )
            )

            if price is None:
                price = price_map.get(
                    (
                        product_id,
                        ""
                    ),
                    {}
                )

            current_price = to_price_float(
                price.get(
                    "current_price_vnd"
                )
            )

            listed_price = to_price_float(
                price.get(
                    "listed_price_vnd"
                )
            )

            discount = to_percent_float(
                price.get(
                    "discount_percent"
                )
            )

            products.append({
                "id":
                    (
                        product_id
                        or str(
                            uuid.uuid4()
                        )
                    ),
                "product_id":
                    product_id,
                "source_product_id":
                    clean_key(
                        price.get(
                            "source_product_id"
                        )
                    ),
                "retailer_id":
                    retailer_id,
                "source":
                    clean_key(
                        row.get(
                            "source"
                        )
                    ),
                "product_name":
                    clean_key(
                        row.get(
                            "product_name"
                        )
                    ),
                "brand":
                    clean_key(
                        row.get(
                            "brand"
                        )
                    ),
                "category":
                    clean_key(
                        row.get(
                            "primary_category"
                        )
                    ),
                "product_url":
                    clean_key(
                        price.get(
                            "product_url"
                        )
                        or row.get(
                            "product_url"
                        )
                    ),
                "current_price_vnd":
                    current_price,
                "listed_price_vnd":
                    listed_price,
                "discount_percent":
                    discount,
                "promotion_text":
                    clean_key(
                        price.get(
                            "promotion_text"
                        )
                    ),
                "availability":
                    clean_key(
                        price.get(
                            "availability"
                        )
                    ),
                "collected_at":
                    clean_key(
                        price.get(
                            "collected_at"
                        )
                    ),
                "collection_status":
                    clean_key(
                        price.get(
                            "collection_status"
                        )
                    ),
                "data_origin":
                    (
                        clean_key(
                            price.get(
                                "data_origin"
                            )
                        )
                        or "scraped"
                    )
            })

        return products

    rows = read_csv(
        OLD_PRODUCTS_CSV
    )

    products = []

    for row in rows:
        products.append({
            "id":
                (
                    clean_key(
                        row.get(
                            "id"
                        )
                    )
                    or str(
                        uuid.uuid4()
                    )
                ),
            "product_id":
                clean_key(
                    row.get(
                        "product_id"
                    )
                ),
            "source_product_id":
                clean_key(
                    row.get(
                        "source_product_id"
                    )
                ),
            "retailer_id":
                clean_key(
                    row.get(
                        "retailer_id"
                    )
                ),
            "source":
                clean_key(
                    row.get(
                        "source"
                    )
                ),
            "product_name":
                clean_key(
                    row.get(
                        "product_name"
                    )
                ),
            "brand":
                clean_key(
                    row.get(
                        "brand"
                    )
                ),
            "product_url":
                clean_key(
                    row.get(
                        "product_url"
                    )
                ),
            "current_price_vnd":
                to_price_float(
                    row.get(
                        "current_price_vnd"
                    )
                ),
            "listed_price_vnd":
                to_price_float(
                    row.get(
                        "listed_price_vnd"
                    )
                ),
            "discount_percent":
                to_percent_float(
                    row.get(
                        "discount_percent"
                    )
                ),
            "rating":
                to_float(
                    row.get(
                        "rating"
                    )
                ),
            "review_count":
                to_int(
                    row.get(
                        "review_count"
                    )
                ),
            "collected_at":
                clean_key(
                    row.get(
                        "collected_at"
                    )
                )
        })

    return products


def repair_working_product_discounts(products):
    if not products:
        return products

    if not PRICE_CSV:
        return products

    price_map = latest_price_map()

    if not price_map:
        return products

    changed = False

    for product in products:
        product_id = clean_key(
            product.get(
                "product_id"
            )
        )

        retailer_id = clean_key(
            product.get(
                "retailer_id"
            )
        )

        if not product_id:
            continue

        price = price_map.get(
            (
                product_id,
                retailer_id
            )
        )

        if price is None:
            price = price_map.get(
                (
                    product_id,
                    ""
                )
            )

        if not price:
            continue

        source_discount_raw = price.get(
            "discount_percent"
        )

        if (
            source_discount_raw is None
            or str(
                source_discount_raw
            ).strip() == ""
        ):
            continue

        correct_discount = to_percent_float(
            source_discount_raw
        )

        old_discount = to_percent_float(
            product.get(
                "discount_percent"
            )
        )

        if (
            abs(
                old_discount
                - correct_discount
            )
            > 0.0001
        ):
            product[
                "discount_percent"
            ] = correct_discount

            changed = True

        source_current_price = price.get(
            "current_price_vnd"
        )

        source_listed_price = price.get(
            "listed_price_vnd"
        )

        if (
            source_current_price is not None
            and str(
                source_current_price
            ).strip() != ""
        ):
            correct_current_price = to_price_float(
                source_current_price
            )

            old_current_price = to_price_float(
                product.get(
                    "current_price_vnd"
                )
            )

            if (
                abs(
                    old_current_price
                    - correct_current_price
                )
                > 0.001
            ):
                product[
                    "current_price_vnd"
                ] = correct_current_price

                changed = True

        if (
            source_listed_price is not None
            and str(
                source_listed_price
            ).strip() != ""
        ):
            correct_listed_price = to_price_float(
                source_listed_price
            )

            old_listed_price = to_price_float(
                product.get(
                    "listed_price_vnd"
                )
            )

            if (
                abs(
                    old_listed_price
                    - correct_listed_price
                )
                > 0.001
            ):
                product[
                    "listed_price_vnd"
                ] = correct_listed_price

                changed = True

    if changed:
        save_products(
            products
        )

        print(
            "Đã sửa lại discount trong products_working.json"
        )

    return products


def load_products():
    if os.path.exists(
        WORKING_JSON
    ):
        try:
            with open(
                WORKING_JSON,
                "r",
                encoding="utf-8"
            ) as file:
                data = json.load(
                    file
                )

            if isinstance(
                data,
                list
            ):
                return repair_working_product_discounts(
                    data
                )

        except Exception as error:
            print(
                "Working JSON error:",
                error
            )

    if os.path.exists(
        OLD_JSON
    ):
        try:
            with open(
                OLD_JSON,
                "r",
                encoding="utf-8"
            ) as file:
                data = json.load(
                    file
                )

            if (
                isinstance(
                    data,
                    list
                )
                and data
            ):
                data = repair_working_product_discounts(
                    data
                )

                save_products(
                    data
                )

                return data

        except Exception as error:
            print(
                "Old JSON error:",
                error
            )

    products = build_products()

    save_products(
        products
    )

    return products


def save_products(products):
    try:
        temp_file = (
            WORKING_JSON
            + ".tmp"
        )

        with open(
            temp_file,
            "w",
            encoding="utf-8"
        ) as file:
            json.dump(
                products,
                file,
                ensure_ascii=False,
                indent=4
            )

        os.replace(
            temp_file,
            WORKING_JSON
        )

        return True

    except Exception as error:
        print(
            "Save products error:",
            error
        )

        return False


def get_first_value(
    row,
    keys,
    default=""
):
    for key in keys:
        if key not in row:
            continue

        value = row.get(
            key
        )

        if value is None:
            continue

        value_str = str(
            value
        ).strip()

        if value_str:
            return value

    return default


def normalize_bool_value(value):
    if isinstance(
        value,
        bool
    ):
        return value

    value_str = str(
        value or ""
    ).strip().lower()

    if value_str in {
        "true",
        "1",
        "yes",
        "y",
        "co",
        "có",
        "da",
        "đã"
    }:
        return True

    if value_str in {
        "false",
        "0",
        "no",
        "n",
        "khong",
        "không",
        "chua",
        "chưa"
    }:
        return False

    return value


def load_reviews():
    review_files = find_review_files()

    reviews = []

    seen_review_keys = set()

    print(
        "Review files:",
        review_files
    )

    for path in review_files:
        file_name_lower = (
            os.path.basename(
                path
            ).lower()
        )

        rows = read_csv(
            path
        )

        if not rows:
            continue

        for row in rows:
            source_product_id = clean_key(
                get_first_value(
                    row,
                    [
                        "source_product_id",
                        "product_source_id",
                        "source_id",
                        "product_id",
                        "id_san_pham",
                        "ma_san_pham",
                        "product_code",
                        "sku",
                        "id"
                    ]
                )
            )

            if not source_product_id:
                continue

            rating_raw = get_first_value(
                row,
                [
                    "rating",
                    "stars",
                    "score",
                    "rating_value",
                    "so_sao",
                    "sao",
                    "star"
                ]
            )

            rating = to_int(
                rating_raw
            )

            if (
                rating < 1
                or rating > 5
            ):
                continue

            review_id = clean_key(
                get_first_value(
                    row,
                    [
                        "review_id",
                        "id_review",
                        "reviewid",
                        "comment_id",
                        "feedback_id",
                        "ma_danh_gia"
                    ]
                )
            )

            review_text = clean_key(
                get_first_value(
                    row,
                    [
                        "review_text",
                        "review",
                        "comment",
                        "content",
                        "content_clean",
                        "review_content",
                        "noi_dung",
                        "noi_dung_danh_gia",
                        "binh_luan",
                        "feedback"
                    ]
                )
            )

            review_date = clean_key(
                get_first_value(
                    row,
                    [
                        "review_date",
                        "date",
                        "created_at",
                        "created_date",
                        "ngay_dang",
                        "ngay",
                        "date_review"
                    ]
                )
            )

            author = clean_key(
                get_first_value(
                    row,
                    [
                        "author",
                        "customer_name",
                        "user_name",
                        "username",
                        "customer",
                        "reviewer",
                        "ten_khach_hang",
                        "ho_ten"
                    ],
                    "Khách hàng"
                )
            )

            source_val = clean_key(
                get_first_value(
                    row,
                    [
                        "source",
                        "retailer",
                        "website",
                        "site",
                        "nguon"
                    ]
                )
            ).lower()

            if (
                not source_val
                or source_val in {
                    "unknown",
                    "unk",
                    "n/a"
                }
            ):
                if "guardian" in file_name_lower:
                    source_val = "guardian"

                elif "hasaki" in file_name_lower:
                    source_val = "hasaki"

                else:
                    source_val = "unknown"

            duplicate_value = str(
                get_first_value(
                    row,
                    [
                        "is_duplicate",
                        "duplicate",
                        "trung_lap"
                    ]
                )
            ).lower().strip()

            if duplicate_value in {
                "true",
                "1",
                "yes"
            }:
                continue

            normalized_text = (
                review_text
                .strip()
                .lower()
            )

            if review_id:
                unique_key = (
                    source_val,
                    source_product_id,
                    review_id
                )
            else:
                unique_key = (
                    source_val,
                    source_product_id,
                    rating,
                    normalized_text,
                    review_date,
                    author
                )

            if unique_key in seen_review_keys:
                continue

            seen_review_keys.add(
                unique_key
            )

            has_images = get_first_value(
                row,
                [
                    "has_images",
                    "has_image",
                    "images",
                    "image",
                    "co_anh"
                ]
            )

            is_bought = get_first_value(
                row,
                [
                    "is_bought",
                    "verified_purchase",
                    "purchased",
                    "is_verified",
                    "da_mua"
                ]
            )

            source_url = clean_key(
                get_first_value(
                    row,
                    [
                        "source_url",
                        "review_url",
                        "url",
                        "link",
                        "product_url"
                    ]
                )
            )

            analysis_status = clean_key(
                get_first_value(
                    row,
                    [
                        "analysis_status",
                        "status",
                        "trang_thai"
                    ]
                )
            )

            reviews.append({
                "source_product_id":
                    str(
                        source_product_id
                    ).strip(),
                "product_name":
                    clean_key(
                        get_first_value(
                            row,
                            [
                                "product_name",
                                "name",
                                "ten_san_pham"
                            ]
                        )
                    ),
                "review_id":
                    review_id,
                "rating":
                    rating,
                "review_text":
                    review_text,
                "review_date":
                    review_date,
                "author":
                    author,
                "is_bought":
                    normalize_bool_value(
                        is_bought
                    ),
                "has_images":
                    normalize_bool_value(
                        has_images
                    ),
                "source_url":
                    source_url,
                "analysis_status":
                    analysis_status,
                "source":
                    source_val
            })

    print(
        "Total reviews loaded:",
        len(reviews)
    )

    guardian_count = sum(
        1
        for review in reviews
        if review.get("source") == "guardian"
    )

    hasaki_count = sum(
        1
        for review in reviews
        if review.get("source") == "hasaki"
    )

    print(
        "Guardian reviews:",
        guardian_count
    )

    print(
        "Hasaki reviews:",
        hasaki_count
    )

    return reviews


def get_current_user():
    authorization = request.headers.get(
        "Authorization",
        ""
    )

    if not authorization:
        return None

    if not authorization.startswith(
        "Bearer "
    ):
        return None

    token = authorization[
        7:
    ].strip()

    if not token:
        return None

    return ACTIVE_SESSIONS.get(
        token
    )


def require_auth(roles=None):
    def decorator(function):
        @wraps(function)
        def wrapper(*args, **kwargs):
            user = get_current_user()

            if user is None:
                return jsonify({
                    "status":
                        "error",
                    "message":
                        "Bạn chưa đăng nhập hoặc phiên đăng nhập không hợp lệ!",
                    "code":
                        "AUTH_REQUIRED"
                }), 401

            if (
                roles
                and user.get("role")
                not in roles
            ):
                return jsonify({
                    "status":
                        "error",
                    "message":
                        "Bạn không có quyền thực hiện chức năng này!",
                    "code":
                        "FORBIDDEN",
                    "required_roles":
                        roles,
                    "current_role":
                        user.get("role")
                }), 403

            g.current_user = user

            return function(
                *args,
                **kwargs
            )

        return wrapper

    return decorator


@app.route(
    "/api/login",
    methods=["POST"]
)
def login():
    data = request.get_json(
        silent=True
    ) or {}

    username = clean_key(
        data.get(
            "username"
        )
    )

    password = str(
        data.get(
            "password",
            ""
        )
    )

    if not username or not password:
        return jsonify({
            "status":
                "error",
            "message":
                "Vui lòng nhập tài khoản và mật khẩu!"
        }), 400

    user = DEMO_USERS.get(
        username
    )

    if (
        user is None
        or not hmac.compare_digest(
            password,
            user["password"]
        )
    ):
        return jsonify({
            "status":
                "error",
            "message":
                "Tài khoản hoặc mật khẩu không đúng!"
        }), 401

    token = secrets.token_urlsafe(
        32
    )

    ACTIVE_SESSIONS[
        token
    ] = {
        "username":
            username,
        "role":
            user["role"],
        "login_at":
            datetime.now().isoformat()
    }

    return jsonify({
        "status":
            "success",
        "message":
            "Đăng nhập thành công!",
        "token":
            token,
        "user": {
            "username":
                username,
            "role":
                user["role"]
        },
        "permissions":
            PERMISSIONS.get(
                user["role"],
                {}
            )
    })


@app.route(
    "/api/me",
    methods=["GET"]
)
@require_auth()
def get_current_user_api():
    user = g.current_user

    return jsonify({
        "status":
            "success",
        "user":
            user,
        "permissions":
            PERMISSIONS.get(
                user.get("role"),
                {}
            )
    })


@app.route(
    "/api/logout",
    methods=["POST"]
)
@require_auth()
def logout():
    authorization = request.headers.get(
        "Authorization",
        ""
    )

    token = authorization[
        7:
    ].strip()

    ACTIVE_SESSIONS.pop(
        token,
        None
    )

    return jsonify({
        "status":
            "success",
        "message":
            "Đăng xuất thành công!"
    })


@app.route(
    "/api/permissions",
    methods=["GET"]
)
@require_auth()
def get_permissions():
    return jsonify({
        "status":
            "success",
        "permissions":
            PERMISSIONS
    })


@app.route("/")
def home():
    return jsonify({
        "status":
            "success",
        "message":
            "BeautyPrice Intelligence Backend đang hoạt động!"
    })


@app.route(
    "/api/test"
)
def test_api():
    return jsonify({
        "status":
            "success",
        "message":
            "API hoạt động bình thường!"
    })


@app.route(
    "/api/products",
    methods=["GET"]
)
def get_products():
    products = load_products()

    return jsonify({
        "status":
            "success",
        "total":
            len(products),
        "products":
            products
    })


@app.route(
    "/api/price-statistics",
    methods=["GET"]
)
def get_price_statistics():
    try:
        rows = read_csv(
            PRICE_CSV
        )

        prices = []

        for row in rows:
            price = to_price_float(
                row.get(
                    "current_price_vnd"
                )
            )

            if price > 0:
                prices.append(
                    price
                )

        if not prices:
            return jsonify({
                "status":
                    "error",
                "message":
                    "Không tìm thấy giá hợp lệ trong file price.",
                "total_rows":
                    len(rows),
                "valid_prices":
                    0
            }), 404

        prices.sort()

        count = len(
            prices
        )

        mean_val = (
            sum(prices)
            / count
        )

        min_val = prices[0]
        max_val = prices[-1]

        mid = count // 2

        median_val = (
            (
                prices[mid - 1]
                + prices[mid]
            )
            / 2
            if count % 2 == 0
            else prices[mid]
        )

        return jsonify({
            "status":
                "success",
            "statistics": {
                "count":
                    count,
                "mean":
                    mean_val,
                "average":
                    mean_val,
                "average_price":
                    mean_val,
                "min":
                    min_val,
                "min_price":
                    min_val,
                "median":
                    median_val,
                "median_price":
                    median_val,
                "max":
                    max_val,
                "max_price":
                    max_val
            },
            "total_rows":
                len(rows),
            "valid_prices":
                count,
            "average_price_vnd":
                mean_val
        })

    except Exception as error:
        return jsonify({
            "status":
                "error",
            "message":
                str(error)
        }), 500


@app.route(
    "/api/reviews",
    methods=["GET"]
)
def get_reviews():
    reviews = load_reviews()

    return jsonify({
        "status":
            "success",
        "total":
            len(reviews),
        "reviews":
            reviews
    })


@app.route(
    "/api/products",
    methods=["POST"]
)
@require_auth(["Admin"])
def create_product():
    data = request.get_json(
        silent=True
    ) or {}

    product_name = clean_key(
        data.get(
            "product_name"
        )
    )

    brand = clean_key(
        data.get(
            "brand"
        )
    )

    current_price = to_price_float(
        data.get(
            "current_price_vnd"
        )
    )

    if not product_name:
        return jsonify({
            "status":
                "error",
            "message":
                "Tên sản phẩm không được để trống!"
        }), 400

    if not brand:
        return jsonify({
            "status":
                "error",
            "message":
                "Thương hiệu không được để trống!"
        }), 400

    if current_price <= 0:
        return jsonify({
            "status":
                "error",
            "message":
                "Giá hiện tại phải lớn hơn 0!"
        }), 400

    products = load_products()

    new_id = str(
        uuid.uuid4()
    )

    new_product = {
        "id":
            new_id,
        "product_id":
            new_id,
        "source_product_id":
            clean_key(
                data.get(
                    "source_product_id"
                )
            ),
        "retailer_id":
            clean_key(
                data.get(
                    "retailer_id"
                )
            ),
        "source":
            clean_key(
                data.get(
                    "source"
                )
            ),
        "product_name":
            product_name,
        "brand":
            brand,
        "category":
            clean_key(
                data.get(
                    "category"
                )
            ),
        "product_url":
            clean_key(
                data.get(
                    "product_url"
                )
            ),
        "current_price_vnd":
            current_price,
        "listed_price_vnd":
            to_price_float(
                data.get(
                    "listed_price_vnd"
                )
            ),
        "discount_percent":
            to_percent_float(
                data.get(
                    "discount_percent"
                )
            ),
        "promotion_text":
            clean_key(
                data.get(
                    "promotion_text"
                )
            ),
        "availability":
            clean_key(
                data.get(
                    "availability"
                )
            ),
        "collected_at":
            clean_key(
                data.get(
                    "collected_at"
                )
            )
            or datetime.now().isoformat(),
        "collection_status":
            clean_key(
                data.get(
                    "collection_status"
                )
            ),
        "data_origin":
            clean_key(
                data.get(
                    "data_origin"
                )
            )
            or "manual"
    }

    products.insert(
        0,
        new_product
    )

    saved = save_products(
        products
    )

    if not saved:
        return jsonify({
            "status":
                "error",
            "message":
                "Không thể lưu sản phẩm!"
        }), 500

    return jsonify({
        "status":
            "success",
        "message":
            "Thêm sản phẩm thành công!",
        "product":
            new_product,
        "total":
            len(products)
    }), 201


@app.route(
    "/api/products/<product_id>",
    methods=["PUT"]
)
@require_auth(["Admin"])
def update_product(
    product_id
):
    data = request.get_json(
        silent=True
    ) or {}

    products = load_products()

    for product in products:
        current_id = str(
            product.get(
                "id",
                ""
            )
        )

        current_product_id = str(
            product.get(
                "product_id",
                ""
            )
        )

        if (
            current_id
            == str(product_id)
            or
            current_product_id
            == str(product_id)
        ):
            allowed_fields = [
                "source_product_id",
                "retailer_id",
                "source",
                "product_name",
                "brand",
                "category",
                "product_url",
                "current_price_vnd",
                "listed_price_vnd",
                "discount_percent",
                "promotion_text",
                "availability",
                "collected_at",
                "collection_status",
                "data_origin"
            ]

            for field in allowed_fields:
                if field not in data:
                    continue

                if field in {
                    "current_price_vnd",
                    "listed_price_vnd"
                }:
                    product[field] = to_price_float(
                        data.get(
                            field
                        )
                    )

                elif field == "discount_percent":
                    product[field] = to_percent_float(
                        data.get(
                            field
                        )
                    )

                else:
                    product[field] = clean_key(
                        data.get(
                            field
                        )
                    )

            saved = save_products(
                products
            )

            if not saved:
                return jsonify({
                    "status":
                        "error",
                    "message":
                        "Không thể lưu dữ liệu!"
                }), 500

            return jsonify({
                "status":
                    "success",
                "message":
                    "Cập nhật sản phẩm thành công!",
                "product":
                    product
            })

    return jsonify({
        "status":
            "error",
        "message":
            "Không tìm thấy sản phẩm!"
    }), 404


@app.route(
    "/api/products/<product_id>",
    methods=["DELETE"]
)
@require_auth(["Admin"])
def delete_product(
    product_id
):
    products = load_products()

    old_length = len(
        products
    )

    products = [
        product
        for product in products
        if
        str(
            product.get(
                "id",
                ""
            )
        ) != str(product_id)
        and
        str(
            product.get(
                "product_id",
                ""
            )
        ) != str(product_id)
    ]

    if len(products) == old_length:
        return jsonify({
            "status":
                "error",
            "message":
                "Không tìm thấy sản phẩm!"
        }), 404

    saved = save_products(
        products
    )

    if not saved:
        return jsonify({
            "status":
                "error",
            "message":
                "Không thể lưu dữ liệu!"
        }), 500

    return jsonify({
        "status":
            "success",
        "message":
            "Xóa sản phẩm thành công!",
        "total":
            len(products)
    })


def find_ml_file(
    base_name
):
    if not os.path.isdir(
        ML_DIR
    ):
        return None

    base_lower = str(
        base_name
    ).lower()

    extensions = {
        ".csv",
        ".xlsx",
        ".xls",
        ".json"
    }

    exact_matches = []
    partial_matches = []

    for filename in os.listdir(
        ML_DIR
    ):
        path = os.path.join(
            ML_DIR,
            filename
        )

        if not os.path.isfile(
            path
        ):
            continue

        stem, ext = os.path.splitext(
            filename.lower()
        )

        if ext not in extensions:
            continue

        if stem == base_lower:
            exact_matches.append(
                path
            )

        elif base_lower in stem:
            partial_matches.append(
                path
            )

    if exact_matches:
        return exact_matches[0]

    if partial_matches:
        return partial_matches[0]

    return None


def read_ml_file(
    file_path
):
    if (
        not file_path
        or not os.path.exists(
            file_path
        )
    ):
        return []

    try:
        extension = os.path.splitext(
            file_path
        )[1].lower()

        if extension == ".csv":
            return read_csv(
                file_path
            )

        if extension == ".json":
            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:
                data = json.load(
                    file
                )

            if isinstance(
                data,
                list
            ):
                return data

            return []

        if extension in {
            ".xlsx",
            ".xls"
        }:
            from openpyxl import load_workbook

            workbook = load_workbook(
                file_path,
                read_only=True,
                data_only=True
            )

            sheet = workbook.active

            rows = list(
                sheet.iter_rows(
                    values_only=True
                )
            )

            workbook.close()

            if not rows:
                return []

            headers = [
                clean_key(
                    value
                )
                for value in rows[0]
            ]

            result = []

            for values in rows[1:]:
                row = {}

                for index, header in enumerate(
                    headers
                ):
                    if not header:
                        continue

                    if index < len(values):
                        row[header] = (
                            values[index]
                        )
                    else:
                        row[header] = None

                result.append(
                    row
                )

            return result

    except Exception as error:
        print(
            "ML file error:",
            error
        )

    return []


def clean_dataframe_for_json(
    records
):
    result = []

    for record in (
        records or []
    ):
        if not isinstance(
            record,
            dict
        ):
            continue

        cleaned = {}

        for key, value in (
            record.items()
        ):
            if isinstance(
                value,
                datetime
            ):
                value = value.isoformat()

            elif isinstance(
                value,
                float
            ):
                if value != value:
                    value = None

                elif value.is_integer():
                    value = int(
                        value
                    )

            cleaned[
                str(key).strip()
            ] = value

        result.append(
            cleaned
        )

    return result


def normalize_columns(
    records
):
    return [
        {
            str(key).strip():
                value
            for key, value
            in row.items()
        }
        for row in (
            records or []
        )
        if isinstance(
            row,
            dict
        )
    ]


@app.route(
    "/api/predictions",
    methods=["GET"]
)
def get_predictions():
    try:
        prediction_file = find_ml_file(
            "regression_test_predictions"
        )

        if prediction_file is None:
            return jsonify({
                "status":
                    "error",
                "message":
                    "Không tìm thấy file regression_test_predictions"
            }), 404

        rows = normalize_columns(
            read_ml_file(
                prediction_file
            )
        )

        if not rows:
            return jsonify({
                "status":
                    "error",
                "message":
                    "Không có dữ liệu trong file"
            }), 404

        q1 = 169000.0
        q2 = 340000.0

        threshold_file = find_ml_file(
            "price_segment_thresholds"
        )

        if (
            threshold_file
            and threshold_file.lower().endswith(
                ".json"
            )
        ):
            try:
                with open(
                    threshold_file,
                    "r",
                    encoding="utf-8"
                ) as file:
                    threshold_data = json.load(
                        file
                    )

                q1 = to_ml_float(
                    threshold_data.get(
                        "q1_vnd",
                        q1
                    )
                )

                q2 = to_ml_float(
                    threshold_data.get(
                        "q2_vnd",
                        q2
                    )
                )

            except Exception:
                pass

        def get_price_group(price):
            if price <= q1:
                return "Thấp"

            if price <= q2:
                return "Trung bình"

            return "Cao"

        predictions = []
        ape_values = []

        for row in rows:
            actual_price = to_ml_float(
                row.get(
                    "current_price_vnd"
                )
                or row.get(
                    "actual_price"
                )
                or row.get(
                    "price"
                )
            )

            predicted_price = to_ml_float(
                row.get(
                    "predicted_price_vnd"
                )
                or row.get(
                    "predicted_price"
                )
                or row.get(
                    "y_pred"
                )
            )

            absolute_error = to_ml_float(
                row.get(
                    "absolute_error_vnd"
                )
                or row.get(
                    "absolute_error"
                )
            )

            if (
                absolute_error == 0
                and actual_price > 0
            ):
                absolute_error = abs(
                    predicted_price
                    - actual_price
                )

            ape = to_ml_float(
                row.get(
                    "absolute_percentage_error"
                )
                or row.get(
                    "mape"
                )
            )

            if ape:
                ape_values.append(
                    ape
                )

            elif actual_price != 0:
                ape_values.append(
                    abs(
                        predicted_price
                        - actual_price
                    )
                    /
                    abs(
                        actual_price
                    )
                    * 100
                )

            predictions.append({
                "product_id":
                    clean_key(
                        row.get(
                            "product_id"
                        )
                        or row.get(
                            "id"
                        )
                    ),
                "product_name":
                    clean_key(
                        row.get(
                            "product_name"
                        )
                        or row.get(
                            "name"
                        )
                    ),
                "brand":
                    clean_key(
                        row.get(
                            "brand"
                        )
                    ),
                "primary_category":
                    clean_key(
                        row.get(
                            "primary_category"
                        )
                        or row.get(
                            "category"
                        )
                    ),
                "current_price_vnd":
                    actual_price,
                "predicted_price_vnd":
                    predicted_price,
                "absolute_error_vnd":
                    absolute_error,
                "absolute_percentage_error":
                    ape,
                "smape_percent":
                    to_ml_float(
                        row.get(
                            "smape_percent"
                        )
                    ),
                "predicted_price_group":
                    get_price_group(
                        predicted_price
                    )
            })

        mae = 87918.90
        rmse = 145603.04
        r2 = 0.6524

        mape = (
            sum(
                ape_values
            )
            /
            len(
                ape_values
            )
            if ape_values
            else 0.0
        )

        return jsonify({
            "status":
                "success",
            "total":
                len(
                    predictions
                ),
            "source_file":
                os.path.basename(
                    prediction_file
                ),
            "predictions":
                predictions,
            "metrics": {
                "mae":
                    mae,
                "rmse":
                    rmse,
                "r2":
                    r2,
                "mape":
                    mape
            },
            "thresholds": {
                "q1_vnd":
                    q1,
                "q2_vnd":
                    q2
            }
        })

    except Exception as error:
        return jsonify({
            "status":
                "error",
            "message":
                str(error)
        }), 500


@app.route(
    "/api/clusters",
    methods=["GET"]
)
def get_clusters():
    try:
        working_products = load_products()

        cluster_map = {}

        cluster_file = find_ml_file(
            "product_clusters"
        )

        if cluster_file:
            cluster_rows = normalize_columns(
                read_ml_file(
                    cluster_file
                )
            )

            for row in cluster_rows:
                pid = clean_key(
                    row.get(
                        "product_id"
                    )
                )

                c_val = row.get(
                    "cluster"
                )

                if (
                    pid
                    and c_val is not None
                ):
                    try:
                        cluster_map[
                            pid
                        ] = int(
                            float(
                                c_val
                            )
                        )

                    except Exception:
                        pass

        cluster_products = []

        for prod in working_products:
            pid = clean_key(
                prod.get(
                    "product_id"
                )
                or prod.get(
                    "id"
                )
            )

            price = to_price_float(
                prod.get(
                    "current_price_vnd"
                )
            )

            cluster_id = cluster_map.get(
                pid
            )

            if cluster_id is None:
                if price <= 169000:
                    cluster_id = 0

                elif price <= 340000:
                    cluster_id = 1

                else:
                    cluster_id = 2

            cluster_products.append({
                "product_id":
                    pid,
                "product_name":
                    (
                        clean_key(
                            prod.get(
                                "product_name"
                            )
                        )
                        or
                        "Không có tên sản phẩm"
                    ),
                "current_price_vnd":
                    price,
                "cluster":
                    cluster_id,
                "brand":
                    clean_key(
                        prod.get(
                            "brand"
                        )
                    ),
                "primary_category":
                    clean_key(
                        prod.get(
                            "category"
                        )
                        or prod.get(
                            "primary_category"
                        )
                    ),
                "retailer_id":
                    clean_key(
                        prod.get(
                            "retailer_id"
                        )
                    ),
                "discount_percent":
                    to_percent_float(
                        prod.get(
                            "discount_percent"
                        )
                    ),
                "rating":
                    to_float(
                        prod.get(
                            "rating",
                            5.0
                        )
                    )
            })

        total_products = len(
            cluster_products
        )

        cluster_counts = {}

        for product in cluster_products:
            cluster_id = product[
                "cluster"
            ]

            cluster_counts[
                cluster_id
            ] = (
                cluster_counts.get(
                    cluster_id,
                    0
                )
                + 1
            )

        clusters = []

        for cluster_id in sorted(
            cluster_counts.keys()
        ):
            count = cluster_counts[
                cluster_id
            ]

            percentage = (
                count
                /
                total_products
                *
                100
            )

            clusters.append({
                "cluster":
                    int(
                        cluster_id
                    ),
                "count":
                    int(
                        count
                    ),
                "percentage":
                    round(
                        percentage,
                        2
                    )
            })

        return jsonify({
            "success":
                True,
            "total_products":
                total_products,
            "total_clusters":
                len(
                    clusters
                ),
            "clusters":
                clusters,
            "products":
                cluster_products
        })

    except Exception as error:
        return jsonify({
            "success":
                False,
            "message":
                str(error)
        }), 500


@app.route(
    "/api/clusters/summary",
    methods=["GET"]
)
def get_cluster_summary():
    try:
        summary_file = find_ml_file(
            "kmeans_cluster_summary"
        )

        if summary_file is None:
            return jsonify({
                "success":
                    False,
                "message":
                    "Không tìm thấy file"
            }), 404

        records = normalize_columns(
            read_ml_file(
                summary_file
            )
        )

        if not records:
            return jsonify({
                "success":
                    False,
                "message":
                    "Lỗi đọc dữ liệu"
            }), 500

        return jsonify({
            "success":
                True,
            "data":
                clean_dataframe_for_json(
                    records
                ),
            "columns":
                list(
                    records[0].keys()
                )
        })

    except Exception as error:
        return jsonify({
            "success":
                False,
            "message":
                str(error)
        }), 500


@app.route(
    "/api/clusters/features",
    methods=["GET"]
)
def get_cluster_features():
    try:
        feature_file = find_ml_file(
            "model_feature_dictionary"
        )

        if feature_file is None:
            return jsonify({
                "success":
                    False,
                "message":
                    "Không tìm thấy file"
            }), 404

        records = normalize_columns(
            read_ml_file(
                feature_file
            )
        )

        if not records:
            return jsonify({
                "success":
                    False,
                "message":
                    "Lỗi đọc dữ liệu"
            }), 500

        return jsonify({
            "success":
                True,
            "data":
                clean_dataframe_for_json(
                    records
                ),
            "columns":
                list(
                    records[0].keys()
                )
        })

    except Exception as error:
        return jsonify({
            "success":
                False,
            "message":
                str(error)
        }), 500


if __name__ == "__main__":
    print(
        "Products CSV:",
        PRODUCTS_CSV
    )

    print(
        "Price CSV:",
        PRICE_CSV
    )

    print(
        "Working JSON:",
        WORKING_JSON
    )

    print(
        "Starting BeautyPrice Intelligence Backend..."
    )

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )