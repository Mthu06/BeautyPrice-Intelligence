"""Tạo bộ từ điển KPI BEAUTYPRICE độc lập với notebook EDA."""
from pathlib import Path
import pandas as pd
from stage6_data import read_reviews, review_metrics, cross_source_gaps, GAP_REASON, MATCH_REASON
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "datar"
OUTPUT_FILE = HERE / "02_kpi_dictionary.xlsx"

def source(pattern):
    files = sorted(DATA.glob(pattern))
    if len(files) != 1: raise FileNotFoundError(f"Cần đúng một file khớp {pattern}")
    return files[0]
def read(pattern): return pd.read_csv(source(pattern), low_memory=False, encoding="utf-8-sig")

snapshots, products, matching = read("price_snapshots*.csv"), read("products*.csv"), read("product_matching.csv")
price = pd.to_numeric(snapshots["current_price_vnd"], errors="coerce").replace([float("inf"),-float("inf")], float("nan"))
listed = pd.to_numeric(snapshots["listed_price_vnd"], errors="coerce").replace([float("inf"),-float("inf")], float("nan"))
valid = price.gt(0); anomalous = valid & listed.gt(0) & listed.lt(price)
pair = valid & listed.gt(0) & listed.ge(price)
dates = pd.to_datetime(snapshots["collected_at"], errors="coerce", utc=True).dt.tz_convert("Asia/Ho_Chi_Minh")
period = f"{dates.min():%Y-%m-%d} đến {dates.max():%Y-%m-%d} (UTC+7); toàn bộ snapshot"
verified = matching["match_status"].astype(str).str.lower().eq("matched") & matching["verified_by_human"].astype(str).str.lower().eq("true")
num = int(verified.sum())
# The previous candidate filter was an implementation choice, not an approved methodology.
# Do not treat it as N_need_matching. A future formal definition must be documented first.
den = None
recalc = (listed[pair]-price[pair])/listed[pair]
gaps=cross_source_gaps(snapshots,matching)

headers = ["kpi_id","kpi_name","business_definition","formula","unit","filter_conditions","time_scope","current_value","status","data_source","validation_method"]
rows = [
 ["KPI13","Tổng sản phẩm duy nhất","Sản phẩm duy nhất trong bảng chuẩn hóa","COUNT DISTINCT(product_id)","sản phẩm","product_id không trống","Toàn bộ products hiện có",products["product_id"].dropna().nunique(),"CALCULATED",source("products*.csv").name,"Đếm độc lập product_id duy nhất"],
 ["KPI14","Tổng bản ghi giá","Tổng snapshot trong phạm vi","COUNT(*)","bản ghi","Không loại theo hợp lệ giá",period,len(snapshots),"CALCULATED",source("price_snapshots*.csv").name,"Đếm số dòng snapshot"],
 ["KPI02","Giá thấp nhất","Giá hiện tại nhỏ nhất hợp lệ","MIN(current_price_vnd)","VND","current_price_vnd hữu hạn > 0",period,price[valid].min(),"CALCULATED",source("price_snapshots*.csv").name,"Đối chiếu MIN trên tập hợp lệ"],
 ["KPI03","Giá cao nhất","Giá hiện tại lớn nhất hợp lệ","MAX(current_price_vnd)","VND","current_price_vnd hữu hạn > 0",period,price[valid].max(),"CALCULATED",source("price_snapshots*.csv").name,"Đối chiếu MAX trên tập hợp lệ"],
 ["KPI04","Giá trung vị","Trung vị giá hiện tại hợp lệ","MEDIAN(current_price_vnd)","VND","current_price_vnd hữu hạn > 0",period,price[valid].median(),"CALCULATED",source("price_snapshots*.csv").name,"Đối chiếu MEDIAN"],
 ["KPI08","Giá trung bình","Trung bình giá hiện tại hợp lệ","SUM(P_i)/n","VND","current_price_vnd hữu hạn > 0",period,price[valid].mean(),"CALCULATED",source("price_snapshots*.csv").name,"Đối chiếu SUM/n"],
 ["KPI12","Cross-source Price Gap","Mean gap across verified equivalent product-date groups","Per verified group/date: MAX(P across retailers)-MIN(P across retailers); card=MEAN(group gaps)","VND","matched + human verified; match_group_id; >=2 retailers; same local date, latest valid snapshot per retailer",period,float(gaps.price_gap_vnd.mean()) if len(gaps) else None,"CALCULATED" if len(gaps) else "NOT CALCULABLE","product_matching.csv + "+source("price_snapshots*.csv").name,"Independent min/max within each verified equivalent group"],
 ["KPI05","Tỷ lệ bản ghi đủ dữ liệu giá","Tỷ lệ snapshot có giá hiện tại hợp lệ","N(current_price_vnd>0)/N(snapshot)","%","Mẫu số là mọi snapshot trong phạm vi",period,valid.mean(),"CALCULATED",source("price_snapshots*.csv").name,"Đếm tử và mẫu độc lập"],
 ["KPI06","Successful Matching Rate","Verified successful matches / formally defined records requiring matching","N_verified_successful_matches / N_need_matching","%","N_need_matching has not been formally defined; no candidate proxy denominator","Trạng thái file matching hiện tại",None,"NOT CALCULABLE",source("product_matching.csv").name,MATCH_REASON],
 ["KPI09","Tỷ lệ giảm giá tính lại trung bình","Trung bình tỷ lệ giảm tính lại từng snapshot","AVG((L_i-P_i)/L_i)","%","current_price_vnd>0 và listed_price_vnd>0",period,recalc.mean() if len(recalc) else None,"CALCULATED" if len(recalc) else "NOT_CALCULABLE",source("price_snapshots*.csv").name,f"Tính lại {len(recalc)} cặp giá"],
]

for row in rows:
 if row[0]=='KPI09':
  row[5]='Finite P>0 and L>0 and L>=P; anomalous L<P excluded'
  row[10]=f'Tính lại {len(recalc)} cặp giá; anomalous_discount={int(anomalous.sum())}'
headers.append('reason')
for row in rows:
 row.append(GAP_REASON if row[0]=='KPI12' and not len(gaps) else MATCH_REASON if row[0]=='KPI06' else '')

wb=Workbook(); ws=wb.active; ws.title="KPI_dictionary"; ws.sheet_view.showGridLines=False
ws["A1"]="BEAUTYPRICE INTELLIGENCE - Bộ từ điển KPI"; ws["A1"].font=Font(name="Arial",size=15,bold=True,color="17365D")
ws["A2"]="KPI được chốt độc lập sau EDA; tỷ lệ lưu dạng 0..1 và hiển thị %."
ws.merge_cells("A1:L1"); ws.merge_cells("A2:L2")
ws.row_dimensions[1].height=25; ws.row_dimensions[2].height=24
for j,h in enumerate(headers,1):
 c=ws.cell(4,j,h); c.font=Font(name="Arial",bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78"); c.alignment=Alignment(wrap_text=True)
for i,row in enumerate(rows,5):
 for j,v in enumerate(row,1): ws.cell(i,j,v)
 if row[4]=="%": ws.cell(i,8).number_format="0.00%"
 elif row[4]=="VND": ws.cell(i,8).number_format="#,##0"
for j,w in enumerate([12,30,48,42,12,72,44,20,22,28,52,65],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:L{len(rows)+4}"
for row in ws.iter_rows():
 for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
ws["A1"].alignment=Alignment(vertical="center",horizontal="left")
ws["A2"].alignment=Alignment(vertical="center",horizontal="left")
for i in range(5,len(rows)+5): ws.row_dimensions[i].height=70
review_paths,reviews=read_reviews()
support=review_metrics(reviews)
review_support_rows=[{'source':'all',**support}]+[{'source':source,**review_metrics(frame)} for source,frame in reviews.groupby('source')]
rs=wb.create_sheet('Review_Metrics')
rs.append(['Supporting Review Metrics – Hasaki & Guardian'])
rs.merge_cells('A1:K1');rs.row_dimensions[1].height=25
rs['A1'].font=Font(name='Arial',size=14,bold=True,color='17365D')
rs.append(['kpi_id','kpi_name','business_definition','formula','unit','filter_conditions','time_scope','current_value','status','data_source','validation_method'])
definitions=[
 ('total','Total reviews','COUNT(*)','reviews','All selected source rows'),
 ('products','Unique reviewed products','COUNT DISTINCT(source, source_product_id)','products','Nonblank source_product_id; not verified price product IDs'),
 ('mean','Average rating','MEAN(rating)','stars','Finite rating in [1,5]'),
 ('median','Median rating','MEDIAN(rating)','stars','Finite rating in [1,5]'),
 ('purchased','Purchased review rate','N(is_bought in yes,true,1)/N(nonblank is_bought)','%','Known is_bought rows; source-reported flag'),
 ('images','Review with image rate','N(has_images=true)/N(nonblank has_images)','%','Known has_images rows'),
 ('unknown_rate','Source product-match flag: unknown rate','N(source product_id_match=unknown)/N(reviews)','%','All selected source reviews; this source field is NOT a verified link to price products'),
]
for group in review_support_rows:
 for key,name,formula,unit,filters in definitions:
  value=group[key]
  rs.append(['REV_'+group['source']+'_'+key,name,'Hasaki & Guardian; independent of price metrics',formula,unit,'source='+group['source']+'; '+filters,'Source review_date; not price collection dates',value,'CALCULATED' if value is not None else 'NOT CALCULABLE','; '.join(str(p.relative_to(DATA)) for p in review_paths),'Reconcile independently from source-filtered dataframe'])
  if unit=='%': rs.cell(rs.max_row,8).number_format='0.00%'
  else: rs.cell(rs.max_row,8).number_format='#,##0.00' if key in ['mean','median'] else '#,##0'
for row in rs:
 for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
for j,w in enumerate([24,32,48,48,14,65,65,20,24,65,65],1): rs.column_dimensions[get_column_letter(j)].width=w
for c in rs[2]: c.fill=PatternFill('solid',fgColor='1F4E78'); c.font=Font(bold=True,color='FFFFFF')
for i in range(3,rs.max_row+1): rs.row_dimensions[i].height=55
rs.freeze_panes='A3'
gs=wb.create_sheet('Verified_Product_Gaps'); gs.append(list(gaps.columns))
for row in gaps.itertuples(index=False,name=None): gs.append(list(row))
if gaps.empty: gs.append([GAP_REASON])
wb.save(OUTPUT_FILE)
print(f"Created: {OUTPUT_FILE}")
