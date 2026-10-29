"""Tạo workbook biểu đồ EDA từ dữ liệu thật BEAUTYPRICE."""
from pathlib import Path
import argparse
import csv
import hashlib
import html
import json
import math
import shutil
import statistics
import subprocess
import sys
from datetime import datetime, timedelta, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import pandas as pd
from stage6_data import read_reviews, review_metrics, GAP_REASON, MATCH_REASON
from stage6_reviews import REVIEW_ENGINE, REVIEW_UI, review_panel, add_review_charts, review_test_cases, review_engine_tests
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE.parent / "datar"
OUTPUT_FILE = HERE / "02_ket_qua_bieu_do.xlsx"


def find_csv(pattern):
    files = sorted(DATA_DIR.glob(pattern))
    if not files:
        raise FileNotFoundError(f"Không tìm thấy file: {pattern}")
    if len(files) != 1:
        raise ValueError(f"Nhiều file khớp {pattern}; cần chọn rõ nguồn")
    return files[0]


snapshots = pd.read_csv(find_csv("price_snapshots*.csv"), low_memory=False, encoding="utf-8-sig")
matching = pd.read_csv(find_csv("product_matching.csv"), low_memory=False, encoding="utf-8-sig")
raw_snapshots = snapshots.copy()

for column in ["current_price_vnd", "listed_price_vnd"]:
    snapshots[column] = pd.to_numeric(snapshots[column], errors="coerce").replace([float("inf"), -float("inf")], float("nan"))
snapshots["collected_at"] = pd.to_datetime(snapshots["collected_at"], errors="coerce", utc=True).dt.tz_convert("Asia/Ho_Chi_Minh")
all_snapshots = snapshots.copy()
snapshots = snapshots.loc[snapshots["current_price_vnd"].gt(0)].copy()
snapshots["category_display"] = snapshots["primary_category"].fillna("Chưa phân loại")
snapshots["observation_date"] = snapshots["collected_at"].dt.date

anomalous_discount = snapshots['listed_price_vnd'].gt(0) & snapshots['listed_price_vnd'].lt(snapshots['current_price_vnd'])
valid_pair = snapshots["listed_price_vnd"].gt(0) & ~anomalous_discount
snapshots["discount_recalculated"] = pd.NA
snapshots.loc[valid_pair, "discount_recalculated"] = (
    snapshots.loc[valid_pair, "listed_price_vnd"] - snapshots.loc[valid_pair, "current_price_vnd"]
) / snapshots.loc[valid_pair, "listed_price_vnd"]
snapshots["discount_recalculated"] = pd.to_numeric(snapshots["discount_recalculated"], errors="coerce")

verified = (
    matching["match_status"].astype(str).str.lower().eq("matched")
    & matching["verified_by_human"].astype(str).str.lower().eq("true")
)

bins = [-1, 100000, 200000, 300000, 400000, 500000, 750000, 1000000, float("inf")]
labels = ["<=100k", "100-200k", "200-300k", "300-400k", "400-500k", "500-750k", "750k-1m", ">1m"]
price_distribution = pd.cut(snapshots["current_price_vnd"], bins=bins, labels=labels).value_counts(sort=False)

category_price = snapshots.groupby("category_display", as_index=False).agg(
    median_price_vnd=("current_price_vnd", "median"), observations=("snapshot_id", "count")
).sort_values("median_price_vnd", ascending=False)

retailer_price = snapshots.groupby("retailer_name", as_index=False).agg(
    median_price_vnd=("current_price_vnd", "median"), observations=("snapshot_id", "count")
).sort_values("median_price_vnd", ascending=False)

discount_retailer = snapshots.dropna(subset=["discount_recalculated"]).groupby("retailer_name", as_index=False).agg(
    average_discount=("discount_recalculated", "mean"), valid_pairs=("snapshot_id", "count")
).sort_values("average_discount", ascending=False)
discount_retailer["average_discount_percent"] = discount_retailer.pop("average_discount") * 100

discount_category = snapshots.dropna(subset=["discount_recalculated"]).groupby("category_display", as_index=False).agg(
    average_discount=("discount_recalculated", "mean"), valid_pairs=("snapshot_id", "count")
).sort_values("average_discount", ascending=False)
discount_category["average_discount_percent"] = discount_category.pop("average_discount") * 100

trend = snapshots.pivot_table(
    index="observation_date", columns="retailer_name", values="current_price_vnd", aggfunc="mean"
).sort_index()
trend.index = trend.index.map(str)

quality = all_snapshots.assign(valid=all_snapshots["current_price_vnd"].gt(0).astype(int)).groupby("retailer_name").agg(total=("snapshot_id", "size"), valid=("valid", "sum")).reset_index()
quality["missing_or_invalid"] = quality["total"] - quality["valid"]

wb = Workbook()
dashboard = wb.active
dashboard.title = "Bieu_do_phan_tich"
data_sheet = wb.create_sheet("Du_lieu_bieu_do")
dashboard.sheet_view.showGridLines = False
data_sheet.sheet_view.showGridLines = False

title_font = Font(name="Arial", size=14, bold=True, color="17365D")
header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="1F4E78")
warning_fill = PatternFill("solid", fgColor="FFF2CC")

dashboard.merge_cells("A2:R2")
dashboard["A2"] = "BEAUTYPRICE INTELLIGENCE - Biểu đồ phân tích dữ liệu thật"
dashboard["A2"].font = title_font
dashboard["A4"] = "Phạm vi"
dashboard["B4"] = f"{snapshots['observation_date'].min()} đến {snapshots['observation_date'].max()} | {len(snapshots):,} snapshot | {snapshots['retailer_id'].nunique()} retailer"
dashboard["A5"] = "Giới hạn"
dashboard["B5"] = (
    "Biểu đồ theo retailer mô tả mặt bằng catalog, chưa phải so sánh cùng sản phẩm đã xác minh. "
    "Đường thời gian mô tả đợt thu thập, chưa chứng minh biến động giá cùng sản phẩm."
)
dashboard["A5"].fill = warning_fill
dashboard["B5"].fill = warning_fill
dashboard["B5"].alignment = Alignment(wrap_text=True)
for col in range(1, 19):
    dashboard.column_dimensions[get_column_letter(col)].width = 9
dashboard.merge_cells('B4:R4')
dashboard.merge_cells('B5:R6')
dashboard['B5'].alignment = Alignment(wrap_text=True, vertical='center')
dashboard.row_dimensions[5].height = 18
dashboard.row_dimensions[6].height = 18
dashboard.sheet_view.zoomScale = 85
dashboard.print_options.horizontalCentered = True
dashboard.sheet_properties.pageSetUpPr.fitToPage = True
dashboard.page_setup.orientation = 'landscape'
dashboard.page_setup.paperSize = dashboard.PAPERSIZE_A3
dashboard.page_setup.fitToWidth = 1
dashboard.page_setup.fitToHeight = 2
dashboard.print_area = 'A1:R76'


def write_table(start_row, start_col, title, frame):
    data_sheet.cell(start_row, start_col, title).font = Font(name="Arial", size=11, bold=True, color="17365D")
    header_row = start_row + 1
    for j, column in enumerate(frame.columns, start=start_col):
        cell = data_sheet.cell(header_row, j, str(column))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        data_sheet.row_dimensions[header_row].height = 32
    for i, row in enumerate(frame.itertuples(index=False, name=None), start=header_row + 1):
        for j, value in enumerate(row, start=start_col):
            data_sheet.cell(i, j, value)
    return header_row, header_row + len(frame)


distribution_frame = pd.DataFrame({"Khoảng giá": labels, "Số quan sát": price_distribution.values})
d_header, d_end = write_table(2, 1, "Phân bố giá", distribution_frame)
c_header, c_end = write_table(2, 5, "Giá trung vị theo danh mục", category_price)
r_header, r_end = write_table(2, 9, "Giá trung vị theo retailer", retailer_price)
dr_header, dr_end = write_table(15, 1, "Discount tính lại theo retailer", discount_retailer)
dc_header, dc_end = write_table(15, 5, "Discount tính lại theo danh mục", discount_category)

trend_frame = trend.reset_index()
t_header, t_end = write_table(28, 1, "Giá trung bình trong giai đoạn quan sát", trend_frame)
q_header, q_end = write_table(39, 1, "Chất lượng giá theo retailer", quality.rename(columns={"retailer_name": "Nhà bán lẻ", "total": "Tổng snapshot", "valid": "Giá hợp lệ", "missing_or_invalid": "Giá thiếu/sai"}))

for col in range(1, 14):
    data_sheet.column_dimensions[get_column_letter(col)].width = 22
for row in data_sheet:
    for cell in row:
        if isinstance(cell.value, (int, float)):
            if math.isfinite(cell.value):
                cell.number_format = "#,##0.00" if cell.value != int(cell.value) else "#,##0"
            else:
                cell.value = None
for row in range(dr_header + 1, dr_end + 1):
    data_sheet.cell(row, 3).number_format = "0.00"
for row in range(dc_header + 1, dc_end + 1):
    data_sheet.cell(row, 7).number_format = "0.00"


def base_chart(chart, title, x_title, y_title, width=14.5, height=8.0):
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.width = width
    chart.height = height
    chart.style = 10
    chart.legend = None
    return chart


hist = base_chart(BarChart(), "Phân bố giá hiện tại", "Khoảng giá (VND)", "Số quan sát")
hist.type = "col"
hist.add_data(Reference(data_sheet, min_col=2, min_row=d_header, max_row=d_end), titles_from_data=True)
hist.set_categories(Reference(data_sheet, min_col=1, min_row=d_header + 1, max_row=d_end))
hist.dLbls = DataLabelList(); hist.dLbls.showVal = True
dashboard.add_chart(hist, "A7")

category_chart = base_chart(BarChart(), "Giá trung vị theo danh mục", "Giá trung vị (VND)", "Danh mục")
category_chart.type = "bar"
category_chart.x_axis.title = "Danh mục"
category_chart.y_axis.title = "Giá trung vị (VND)"
category_chart.add_data(Reference(data_sheet, min_col=6, min_row=c_header, max_row=c_end), titles_from_data=True)
category_chart.set_categories(Reference(data_sheet, min_col=5, min_row=c_header + 1, max_row=c_end))
dashboard.add_chart(category_chart, "J7")

retailer_chart = base_chart(BarChart(), "Giá trung vị theo retailer (mô tả catalog)", "Retailer", "Giá trung vị (VND)")
retailer_chart.type = "col"
retailer_chart.add_data(Reference(data_sheet, min_col=10, min_row=r_header, max_row=r_end), titles_from_data=True)
retailer_chart.set_categories(Reference(data_sheet, min_col=9, min_row=r_header + 1, max_row=r_end))
dashboard.add_chart(retailer_chart, "A23")

trend_chart = base_chart(LineChart(), "Giá trung bình trong giai đoạn quan sát", "Ngày quan sát", "Giá trung bình (VND)")
trend_chart.add_data(Reference(data_sheet, min_col=2, max_col=1 + len(trend.columns), min_row=t_header, max_row=t_end), titles_from_data=True)
trend_chart.set_categories(Reference(data_sheet, min_col=1, min_row=t_header + 1, max_row=t_end))
for series, retailer_name in zip(trend_chart.series, trend.columns):
    series.tx = SeriesLabel(v=str(retailer_name))
trend_chart.legend = Legend(legendPos="b")
trend_chart.marker = True
for series in trend_chart.series:
    series.marker.symbol = "circle"
    series.marker.size = 5
dashboard.add_chart(trend_chart, "J23")

quality_chart = base_chart(BarChart(), "Số snapshot có giá hợp lệ và thiếu/sai theo retailer", "Số snapshot", "Retailer")
quality_chart.type = "bar"
quality_chart.x_axis.title = "Retailer"
quality_chart.y_axis.title = "Số snapshot"
quality_chart.grouping = "stacked"
quality_chart.overlap = 100
quality_chart.add_data(Reference(data_sheet, min_col=3, max_col=4, min_row=q_header, max_row=q_end), titles_from_data=True)
quality_chart.set_categories(Reference(data_sheet, min_col=1, min_row=q_header + 1, max_row=q_end))
quality_chart.legend = Legend(legendPos="b")
for series, name in zip(quality_chart.series, ["Giá hợp lệ", "Giá thiếu/sai"]):
    series.tx = SeriesLabel(v=name)
dashboard.add_chart(quality_chart, "A39")
dashboard["A55"] = "Lưu ý diễn giải"
dashboard["A55"].font = Font(name="Arial", size=11, bold=True, color="17365D")
dashboard["A56"] = (
    "Các biểu đồ theo retailer là thống kê mô tả catalog quan sát được. Cơ cấu sản phẩm khác nhau giữa retailer, "
    "nên không dùng để kết luận retailer nào rẻ nhất. Biểu đồ thời gian chỉ mô tả giai đoạn quan sát, không phải xu hướng dài hạn."
)
dashboard["A56"].alignment = Alignment(wrap_text=True, vertical="top")
dashboard.merge_cells("A56:R59")
dashboard.print_area = 'A1:R59'

# Insight: mỗi nhận định gắn với số liệu tính từ dữ liệu thật và một bảng chứng cứ.
total = len(all_snapshots)
p = snapshots["current_price_vnd"]
q1, q3 = p.quantile(0.25), p.quantile(0.75)
iqr = q3 - q1
outliers = int(((p < q1 - 1.5 * iqr) | (p > q3 + 1.5 * iqr)).sum())
pair_count = int(valid_pair.sum())
named_categories = category_price.loc[category_price["category_display"].ne("Chưa phân loại")]
cat_hi, cat_lo = named_categories.iloc[0], named_categories.iloc[-1]
disc_hi, disc_lo = discount_retailer.iloc[0], discount_retailer.iloc[-1]
missing_category = int(all_snapshots["primary_category"].isna().sum())
period_label = f"{snapshots['observation_date'].min()} đến {snapshots['observation_date'].max()} (UTC+7)"
insight_rows = [
    ["I01", "Giá tập trung nhiều nhất ở khoảng trên 100.000 đến 200.000 VND", f"{int(price_distribution.iloc[1]):,}/{len(p):,} snapshot giá hợp lệ ({price_distribution.iloc[1]/len(p):.2%}), {period_label}.", "Khoảng này là vùng giá phổ biến nhất trong mẫu quan sát; mẫu gồm nhiều thương hiệu và quy cách.", "Ưu tiên bộ lọc khoảng giá và xem riêng từng danh mục trước khi đặt ngưỡng theo dõi.", f"Du_lieu_bieu_do!A{d_header}:B{d_end}"],
    ["I02", "Đuôi giá cao làm giá trung bình cao hơn trung vị", f"n={len(p):,}; trung bình={p.mean():,.2f} VND; trung vị={p.median():,.0f} VND; IQR={iqr:,.0f} VND; {outliers} ứng viên vượt ngưỡng IQR ({outliers/len(p):.2%}).", "Giá cao có thể liên quan phân khúc, dung tích và quy cách, chưa đủ kết luận là lỗi.", "Ưu tiên trung vị khi mô tả mặt bằng giá; rà soát ứng viên ngoại lệ trên nguồn và giữ nguyên dữ liệu.", "01_ket_qua_eda.xlsx / 02_Numeric_Descriptive và 12_Price_Outlier_Candidates"],
    ["I03", "Mặt bằng giá trung vị khác nhau giữa danh mục", f"{cat_hi['category_display']}: {cat_hi['median_price_vnd']:,.0f} VND (n={int(cat_hi['observations'])}); {cat_lo['category_display']}: {cat_lo['median_price_vnd']:,.0f} VND (n={int(cat_lo['observations'])}); {period_label}.", "Khác biệt phản ánh cơ cấu sản phẩm của các danh mục; không phải chênh lệch của cùng sản phẩm.", "Lọc theo danh mục và thương hiệu khi so sánh mức giá; chỉ so sánh trực tiếp khi đã xác minh tương đương.", f"Du_lieu_bieu_do!E{c_header}:G{c_end}"],
    ["I04", "Mức giảm giá tính lại trung bình khác nhau theo retailer", f"{disc_hi['retailer_name']}: {disc_hi['average_discount_percent']:.2f}% (n={int(disc_hi['valid_pairs'])} cặp giá); {disc_lo['retailer_name']}: {disc_lo['average_discount_percent']:.2f}% (n={int(disc_lo['valid_pairs'])}); chỉ cặp giá dương.", "Độ sâu khuyến mãi quan sát được khác nhau; chưa chứng minh retailer có giá cuối cùng thấp hơn.", "Hiển thị cả giá hiện tại, giá niêm yết và discount tính lại; rà soát discount nguồn lệch giá trị tính lại.", f"Du_lieu_bieu_do!A{dr_header}:C{dr_end}"],
    ["I05", "Giá hiện tại đầy đủ hơn giá niêm yết", f"{len(p):,}/{total:,} giá hiện tại hợp lệ ({len(p)/total:.2%}); {total-pair_count:,}/{total:,} snapshot thiếu cặp giá hợp lệ ({(total-pair_count)/total:.2%}); discount trung bình chỉ dùng {pair_count:,} cặp.", "Không thể tính khuyến mãi cho các bản ghi thiếu giá niêm yết; cần công bố mẫu số thay vì coi là không giảm giá.", "Giữ discount trống khi thiếu cặp giá; ưu tiên bổ sung giá niêm yết từ nguồn thu thập.", f"Du_lieu_bieu_do!A{q_header}:D{q_end}; KPI giảm giá / Dau_vet_gia"],
    ["I06", "Một phần snapshot chưa có primary_category", f"{missing_category:,}/{total:,} snapshot thiếu primary_category ({missing_category/total:.2%}); nhóm này được giữ là Chưa phân loại.", "Tổng hợp chỉ trên danh mục đã gán sẽ bỏ sót một phần mẫu; quan hệ category nhiều-nhiều cần được phân biệt với primary_category.", "Rà soát quy tắc chọn danh mục chính; giữ nhóm thiếu hiển thị rõ và không tự đoán nhãn để làm đẹp kết quả.", f"Du_lieu_bieu_do!E{c_header}:G{c_end}"],
    ["I07", "Chưa có đối chiếu được xác minh thành công", f"0 dòng matched + verified trong {len(matching)} dòng ứng viên/rà soát; KPI chỉ đo hàng đợi trường hợp product_id+retailer_id hiện có.", "Tên giống hoặc độ tin cậy thuật toán cao chưa đủ chứng minh tương đương về biến thể và quy cách.", "Ưu tiên xác minh tên, thương hiệu, dung tích, quy cách và thời điểm trước khi đưa ra kết luận nơi bán rẻ nhất.", "product_matching.csv; Dinh_nghia_KPI / KPI06"],
    ["I08", "Thời gian quan sát ngắn và cơ cấu catalog có thể thay đổi theo ngày", f"{all_snapshots['collected_at'].dt.date.nunique()} ngày thu thập, {period_label}; {total:,} snapshot trên 5 retailer.", "Đường giá trung bình giữa các ngày có thể bị ảnh hưởng bởi cơ cấu mặt hàng được thu thập, chưa phải lịch sử biến động của cùng sản phẩm.", "Thu thập lặp cùng sản phẩm đã xác minh ở nhiều thời điểm; chỉ mô tả đợt quan sát hiện tại.", f"Du_lieu_bieu_do!A{t_header}:F{t_end}"],
]
for insight in insight_rows:
    if insight[0] == 'I07':
        insight[2] += ' Trạng thái file matching hiện tại, chưa có ngày xác minh riêng.'
        insight[5] = 'product_matching.csv; 02_kpi_dictionary.xlsx / KPI_dictionary / KPI06'
    elif period_label not in insight[2]:
        insight[2] += f' Phạm vi: {period_label}.'
review_paths, reviews = read_reviews()
review_stats = review_metrics(reviews)
# Eight evidence-backed themes; retain correct price insights and replace outdated matching claims.
ret_hi,ret_lo=retailer_price.iloc[0],retailer_price.iloc[-1]
insight_rows=[insight_rows[2],insight_rows[1],
 ['I03','Mặt bằng giá catalog khác nhau giữa nhà bán lẻ',f"{ret_hi['retailer_name']}: trung vị {ret_hi['median_price_vnd']:,.0f} VND (n={int(ret_hi['observations'])}); {ret_lo['retailer_name']}: {ret_lo['median_price_vnd']:,.0f} VND (n={int(ret_lo['observations'])}). Phạm vi {period_label}; giá snapshot hợp lệ.",'Khác biệt về cơ cấu catalog không chứng minh nơi nào rẻ hơn cho cùng một sản phẩm.','Lọc danh mục và thương hiệu để mô tả catalog; chỉ so sánh cùng sản phẩm sau xác minh.','02_ket_qua_bieu_do.xlsx / Du_lieu_bieu_do!I3:K8'],
 insight_rows[3],insight_rows[4],
 ['I06','Chưa đủ cơ sở so sánh cùng sản phẩm giữa nhà bán lẻ',f"{len(matching)} dòng matching; {int(verified.sum())} dòng matched + verified. Mẫu số N_need_matching chưa có phương pháp chính thức trong tài liệu dự án. Trạng thái file matching hiện tại.",'Chưa có verified equivalent product matches nên dataset chưa đủ cơ sở kết luận retailer nào có giá thấp nhất cho CÙNG MỘT SẢN PHẨM. Cross-source gap và matching success rate hiện NOT CALCULABLE.','Xác minh thương hiệu, biến thể, dung tích/quy cách; phê duyệt N_need_matching trước khi công bố tỷ lệ.','02_kpi_dictionary.xlsx / KPI_dictionary / KPI12, KPI06'],
 ['I07','Mẫu review gộp chịu ảnh hưởng mạnh bởi Hasaki',f"Hasaki + Guardian: n={review_stats['total']} review; {review_stats['products']} cặp nguồn + mã sản phẩm; pooled rating trung bình {review_stats['mean']:.4f}/5, trung vị {review_stats['median']:.1f}/5; {review_stats['unknown']}/{review_stats['total']} dòng có source product-match flag = unknown ({review_stats['unknown_rate']:.2%}).",'Pooled review metrics không phải comparison cân bằng giữa hai retailer vì tỷ trọng mẫu lệch về Hasaki. Đây là observed rating distribution in the collected sample; không suy rộng cho toàn bộ khách hàng. Cờ nguồn không phải verified link to price products.','Xem per-source metrics khi mô tả Hasaki và Guardian; không dùng pooled average để kết luận retailer nào có customer satisfaction cao hơn. Giữ review độc lập với price.','01_ket_qua_eda.xlsx / 03_Categorical_Descriptive, 06_Review_Descriptive, 07_Review_Data_Quality'],
 insight_rows[7]]
for i,row in enumerate(insight_rows,1): row[0]=f'I{i:02}'
insight_rows[3][2]+=f' anomalous_discount (L<P): {int(anomalous_discount.sum())} dòng; loại khỏi discount hợp lệ.'
insight_rows[6][2]+=f" Khoảng review_date nguồn: {reviews['review_date_parsed'].min():%Y-%m-%d} đến {reviews['review_date_parsed'].max():%Y-%m-%d}; không lọc theo ngày thu thập giá. " + '; '.join(f"{source}: {len(frame)} review ({len(frame)/len(reviews):.2%}), rating quan sát {frame['rating_numeric'].min():g}–{frame['rating_numeric'].max():g}/5, trung bình {frame['rating_numeric'].mean():.4f}/5; source product-match flag counts: {frame['product_id_match'].fillna('MISSING').value_counts().to_dict()}" for source,frame in reviews.groupby('source')) + '. This source field is NOT a verified link to price products.'
insight_rows[4][5]='01_ket_qua_eda.xlsx / 04_Data_Coverage_Quality; 02_ket_qua_bieu_do.xlsx / Du_lieu_bieu_do!A40:D45'
data_sheet['A48'] = 'Nguồn dữ liệu'
data_sheet['B48'] = find_csv('price_snapshots*.csv').name
data_sheet['A49'] = 'Thời gian UTC+7'
data_sheet['B49'] = period_label
data_sheet['B49'].alignment = Alignment(wrap_text=True)
data_sheet.row_dimensions[49].height = 34

def report_table(sheet, headers, rows, widths):
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        sheet.append([v.item() if hasattr(v, "item") else v for v in row])
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row > 1:
                cell.font = Font(name="Arial", size=10)
        sheet.row_dimensions[row[0].row].height = max(24, max(math.ceil(len(str(c.value or "")) / max(8, widths[c.column-1]*.95)) for c in row)*14+8)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions

insight_path = HERE / '03_insight.md'
insight_text = '# Insight BEAUTYPRICE INTELLIGENCE\n\nPhạm vi: ' + period_label + '. Các nhận định mô tả snapshot trong tập thu thập; chưa phải so sánh cùng sản phẩm đã xác minh.\n\n'
for row in insight_rows:
    insight_text += f'## {row[0]} — {row[1]}\n\nPhát hiện: {row[1]}.\n\nSố liệu: {row[2]}\n\nÝ nghĩa: {row[3]}\n\nHành động: {row[4]}\n\nBằng chứng: `{row[5]}`.\n\n'

# Đọc CSV bằng thư viện chuẩn để đối chiếu độc lập với engine JavaScript.
def csv_records(pattern):
    with find_csv(pattern).open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))

def number(value):
    try:
        v = float(value)
        return v if math.isfinite(v) else None
    except (ValueError, TypeError):
        return None

def day_local(value):
    try:
        return datetime.fromisoformat(value).astimezone(timezone(timedelta(hours=7))).date().isoformat()
    except (ValueError, TypeError):
        return ""

raw_records = csv_records("price_snapshots*.csv")
raw_products = csv_records("products*.csv")
raw_matches = csv_records("product_matching.csv")
master_ids = sorted({r['product_id'] for r in raw_products if r['product_id'].strip()})
queue = {}
for r in raw_matches:
    if r['record_type'] in ('matching_candidate', 'review_only_candidate') and r['product_id'].strip() and r['retailer_id'].strip() and (number(r['candidate_retailer_count']) or 0) >= 2:
        key = (r['product_id'], r['retailer_id'])
        queue[key] = queue.get(key, False) or (r['match_status'].lower() == 'matched' and r['verified_by_human'].lower() == 'true')

product_lookup = {r['product_id']: r for r in raw_products}
rows = [[r['product_id'], r['retailer_id'], r['retailer_name'], r['primary_category'] or 'Chưa phân loại', r['brand'] or 'Chưa rõ thương hiệu', number(r['current_price_vnd']), number(r['listed_price_vnd']), day_local(r['collected_at']), product_lookup.get(r['product_id'], {}).get('product_name', r['product_id']), r.get('product_url', ''),r['collected_at'],r['snapshot_id']] for r in raw_records]
verified_identity={}
for r in raw_matches:
    if r['match_status'].lower()=='matched' and r['verified_by_human'].lower()=='true' and r['match_group_id'].strip():
        identity=r['product_id']+'|'+r['retailer_id']
        if identity in verified_identity and verified_identity[identity]!=r['match_group_id']:
            raise ValueError('Ambiguous verified group assignment')
        verified_identity[identity]=r['match_group_id']
queue_rows = [[key[0], key[1], value] for key, value in queue.items()]
metric_keys = ['products', 'snapshots', 'min', 'max', 'median', 'mean', 'gap', 'coverage', 'matching', 'discount']
metric_names = ['Tổng sản phẩm duy nhất', 'Tổng bản ghi giá', 'Giá thấp nhất', 'Giá cao nhất', 'Giá trung vị', 'Giá trung bình', 'Cross-source Price Gap', 'Tỷ lệ bản ghi đủ dữ liệu giá', 'Successful Matching Rate', 'Tỷ lệ giảm giá tính lại (trung bình)']

def independent(filters):
    def chosen(value, selection):
        return not selection or value in selection if isinstance(selection, list) else not selection or value == selection
    def has_filters():
        return any(bool(v) for v in filters.values())
    selected = []
    for r in raw_records:
        pr = number(r['current_price_vnd'])
        date = day_local(r['collected_at'])
        key = (r['product_id'], r['retailer_id'])
        if not chosen(r['retailer_name'], filters.get('retailer')): continue
        if not chosen(r['primary_category'] or 'Chưa phân loại', filters.get('category')): continue
        if not chosen(r['brand'] or 'Chưa rõ thương hiệu', filters.get('brand')): continue
        if not chosen(r['product_id'], filters.get('product')): continue
        if filters.get('min') not in (None, '') and (pr is None or pr < float(filters['min'])): continue
        if filters.get('max') not in (None, '') and (pr is None or pr > float(filters['max'])): continue
        if filters.get('from') and (not date or date < filters['from']): continue
        if filters.get('to') and (not date or date > filters['to']): continue
        if filters.get('quality') == 'valid' and not (pr is not None and pr > 0): continue
        if filters.get('quality') == 'invalid' and pr is not None and pr > 0: continue
        if filters.get('match') == 'pending' and (key not in queue or queue[key]): continue
        if filters.get('match') == 'verified' and not queue.get(key, False): continue
        selected.append(r)
    prices = sorted(number(r['current_price_vnd']) for r in selected if number(r['current_price_vnd']) is not None and number(r['current_price_vnd']) > 0)
    discounts = [(number(r['listed_price_vnd'])-number(r['current_price_vnd']))/number(r['listed_price_vnd']) for r in selected if number(r['listed_price_vnd']) is not None and number(r['current_price_vnd']) is not None and number(r['listed_price_vnd']) >= number(r['current_price_vnd']) > 0]
    latest={}
    for r in selected:
        group_id=verified_identity.get(r['product_id']+'|'+r['retailer_id'])
        value=number(r['current_price_vnd']); date=day_local(r['collected_at'])
        if group_id and date and value is not None and value>0:
            k=(group_id,date,r['retailer_id'])
            if k not in latest or (r['collected_at'],r['snapshot_id'])>(latest[k]['collected_at'],latest[k]['snapshot_id']): latest[k]=r
    equivalent={}
    for (group_id,date,_),r in latest.items(): equivalent.setdefault((group_id,date),[]).append(number(r['current_price_vnd']))
    gaps=[max(v)-min(v) for v in equivalent.values() if len(v)>=2]
    keys = {(r['product_id'], r['retailer_id']) for r in selected}
    active_queue = {k: v for k, v in queue.items() if not has_filters() or k in keys}
    metrics = dict(zip(metric_keys, [len({r['product_id'] for r in selected} & set(master_ids)) if has_filters() else len(master_ids), len(selected), min(prices) if prices else None, max(prices) if prices else None, statistics.median(prices) if prices else None, statistics.mean(prices) if prices else None, statistics.mean(gaps) if gaps else None, len(prices)/len(selected) if selected else None, None, statistics.mean(discounts) if discounts else None]))
    return {'metrics': metrics, 'checks': {'hist': len(prices), 'category': len(prices), 'discountPairs': len(discounts), 'quality': len(selected)}, 'matching_n': len(active_queue), 'matching_ok': sum(active_queue.values())}

tests = [
    ['T01', 'Mở dashboard mặc định', {}],
    ['T02', 'Lọc danh mục Serum', {'category': 'Serum'}],
    ['T03', 'Lọc retailer Hasaki', {'retailer': 'Hasaki'}],
    ['T04', 'Lọc giá từ 100.000 đến 200.000 VND (gồm hai biên)', {'min': '100000', 'max': '200000'}],
    ['T05', 'Lọc ngày 18–19/10/2026', {'from': '2026-10-18', 'to': '2026-10-19'}],
    ['T06', 'Xóa bộ lọc để trở về mặc định', {}],
    ['T07', 'Đối chiếu tất cả 10 KPI', {}],
    ['T08', 'Khoảng giá không có dữ liệu', {'min': '900000000'}],
    ['T09', 'Lọc thương hiệu Cocoon', {'brand': 'Cocoon'}],
    ['T10', 'Kết hợp Guardian, Sữa tắm, giá tối đa 300.000', {'retailer': 'Guardian', 'category': 'Sữa tắm', 'max': '300000'}],
    ['T11', 'Giá đúng biên 5.000 VND', {'min': '5000', 'max': '5000'}],
    ['T12', 'Chỉ trường hợp đã xác minh (hiện không có)', {'match': 'verified'}],
    ['T13', 'Chỉ giá thiếu/không hợp lệ (hiện không có)', {'quality': 'invalid'}],
    ['T14', 'Khoảng giá đảo ngược cần thông báo', {'min': '200000', 'max': '100000'}],
    ['T15', 'Ngày không có dữ liệu', {'from': '2099-01-01', 'to': '2099-01-02'}],
    ['T16', 'Giữ nhóm chưa phân loại', {'category': 'Chưa phân loại'}],
    ['T17', 'Lọc hàng đợi pending', {'match': 'pending'}],
]
tests += [
    ['T18', 'Chọn đồng thời Hasaki và Guardian', {'retailer': ['Hasaki', 'Guardian']}],
    ['T19', 'Chọn đồng thời Serum và Sữa rửa mặt', {'category': ['Serum', 'Sữa rửa mặt']}],
    ['T20', 'Chọn nhiều thương hiệu Cocoon và EUCERIN', {'brand': ['Cocoon', 'EUCERIN']}],
    ['T21', 'Hai retailer và hai danh mục, OR trong nhóm và AND giữa nhóm', {'retailer': ['Hasaki', 'Guardian'], 'category': ['Serum', 'Sữa rửa mặt']}],
    ['T22', 'Danh sách lựa chọn rỗng nghĩa là tất cả', {'retailer': [], 'category': [], 'brand': []}],
]
test_cases = [{'id': tid, 'name': name, 'filters': f, 'expected': independent(f)} for tid, name, f in tests]
payload = {'rows': rows, 'queue': queue_rows, 'masterIds': master_ids, 'metricKeys': metric_keys, 'metricNames': metric_names, 'period': period_label, 'sourceHashes': {find_csv(p).name: hashlib.sha256(find_csv(p).read_bytes()).hexdigest() for p in ['price_snapshots*.csv', 'products*.csv', 'product_matching.csv']}}
payload.update(verifiedIdentity=verified_identity,matchingDenominatorDefined=False,gapReason=GAP_REASON,matchingReason=MATCH_REASON)
for review_path in review_paths:
    payload['sourceHashes'][str(review_path.relative_to(DATA_DIR))]=hashlib.sha256(review_path.read_bytes()).hexdigest()
payload['reviews']=[{'source':r['source'],'product':r['reviewed_product_key'] if pd.notna(r['reviewed_product_key']) else None,'category':r['category_display'],'rating':float(r['rating_numeric']) if pd.notna(r['rating_numeric']) else None,'bought':str(r['is_bought']).strip().lower() if pd.notna(r['is_bought']) else None,'images':str(r['has_images']).strip().lower() if pd.notna(r['has_images']) else None,'matching':r['product_id_match'] if pd.notna(r['product_id_match']) else None} for _,r in reviews.iterrows()]

ENGINE = r"""
function median(a){if(!a.length)return null;const s=[...a].sort((x,y)=>x-y),m=Math.floor(s.length/2);return s.length%2?s[m]:(s[m-1]+s[m])/2;}
function average(a){return a.length?a.reduce((s,x)=>s+x,0)/a.length:null;}
function key(r){return r[0]+'|'+r[1];}
function active(f){return Object.values(f).some(x=>Array.isArray(x)?x.length>0:x!==''&&x!==null&&x!==undefined);}
function chosen(value,selection){return Array.isArray(selection)?selection.length===0||selection.includes(value):!selection||value===selection;}
function calculate(data,f={}){
 const queue=new Map(data.queue.map(r=>[key(r),r[2]])), master=new Set(data.masterIds);
 const selected=data.rows.filter(r=>{
  if(!chosen(r[2],f.retailer))return false;
  if(!chosen(r[3],f.category))return false;
  if(!chosen(r[4],f.brand))return false;
  if(f.min!==undefined&&f.min!==''&&(r[5]===null||r[5]<Number(f.min)))return false;
  if(f.max!==undefined&&f.max!==''&&(r[5]===null||r[5]>Number(f.max)))return false;
  if(f.from&&(!r[7]||r[7]<f.from))return false;
  if(f.to&&(!r[7]||r[7]>f.to))return false;
  if(f.quality==='valid'&&!(r[5]>0))return false;
  if(f.quality==='invalid'&&r[5]>0)return false;
  if(f.match==='pending'&&(!queue.has(key(r))||queue.get(key(r))))return false;
  if(f.match==='verified'&&!queue.get(key(r)))return false;
  return true;
 });
 const valid=selected.filter(r=>r[5]>0), prices=valid.map(r=>r[5]);
 const pair=valid.filter(r=>r[6]>0&&r[6]>=r[5]), discounts=pair.map(r=>(r[6]-r[5])/r[6]);
 const latest=new Map();for(const r of valid){const g=data.verifiedIdentity[key(r)];if(!g||!r[7])continue;const k=g+'|'+r[7]+'|'+r[1],old=latest.get(k);if(!old||r[10]>old[10]||(r[10]===old[10]&&r[11]>old[11]))latest.set(k,r);}
 const groups=new Map();for(const r of latest.values()){const k=data.verifiedIdentity[key(r)]+'|'+r[7];if(!groups.has(k))groups.set(k,[]);groups.get(k).push(r[5]);}const gaps=[...groups.values()].filter(v=>v.length>=2).map(v=>Math.max(...v)-Math.min(...v));
 const keys=new Set(selected.map(key)), cases=data.queue.filter(r=>!active(f)||keys.has(key(r))), ok=cases.filter(r=>r[2]).length;
 const metrics={products:active(f)?new Set(selected.filter(r=>master.has(r[0])).map(r=>r[0])).size:master.size,snapshots:selected.length,min:prices.length?Math.min(...prices):null,max:prices.length?Math.max(...prices):null,median:median(prices),mean:average(prices),gap:average(gaps),coverage:selected.length?prices.length/selected.length:null,matching:data.matchingDenominatorDefined&&Number.isFinite(data.matchingDenominator)&&data.matchingDenominator>0?ok/data.matchingDenominator:null,discount:average(discounts)};
 return {metrics,checks:{hist:valid.length,category:valid.length,discountPairs:pair.length,quality:selected.length},matching_n:cases.length,matching_ok:ok,selected,valid,pair};
}
function closeValue(a,b){return a===null||b===null?a===b:Math.abs(a-b)<=Math.max(1e-8,Math.abs(b)*1e-10);}
function compare(actual,expected){return Object.keys(expected.metrics).every(k=>closeValue(actual.metrics[k],expected.metrics[k]))&&Object.keys(expected.checks).every(k=>actual.checks[k]===expected.checks[k])&&actual.matching_n===expected.matching_n&&actual.matching_ok===expected.matching_ok;}
"""

UI = r"""
const $=id=>document.getElementById(id), esc=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const vi=new Intl.NumberFormat('vi-VN',{maximumFractionDigits:0}),pct=new Intl.NumberFormat('vi-VN',{minimumFractionDigits:2,maximumFractionDigits:2});
const colors=['#cf7b5d','#537c80','#c6a75d','#7c7098','#7a9e69','#ad6e8a','#8395a8'];
function fmt(k,v){if(v===null)return ['gap','matching'].includes(k)?'NOT CALCULABLE':'Không có dữ liệu';return ['coverage','matching','discount'].includes(k)?pct.format(v*100)+'%':vi.format(v);}
const fields=['min','max','from','to','quality','match'];
const selections={retailer:[],category:[],brand:[]};
const options={};
for(const [id,index] of [['retailer',2],['category',3],['brand',4]])options[id]=[...new Set(DATA.rows.map(r=>r[index]))].sort((a,b)=>a.localeCompare(b,'vi')).map(v=>({value:v,label:v}));
const normalize=s=>s.normalize('NFD').replace(/[\u0300-\u036f]/g,'').toLowerCase().replace(/đ/g,'d');
function drawChoices(id){
 const query=normalize($(id+'Search').value),all=options[id].filter(o=>normalize(o.label+' '+o.value).includes(query)),shown=all.slice(0,id==='brand'?80:all.length);
 $(id).innerHTML=shown.map(o=>`<label class="choice"><input type="checkbox" value="${esc(o.value)}" data-field="${id}" ${selections[id].includes(o.value)?'checked':''}><span>${esc(o.label)}</span></label>`).join('');
 $(id+'Count').textContent=`${shown.length}/${all.length} lựa chọn; gõ để tìm thêm`;
 $(id).querySelectorAll('input').forEach(input=>input.addEventListener('change',()=>{const v=input.value;if(input.checked&&!selections[id].includes(v))selections[id].push(v);else if(!input.checked)selections[id]=selections[id].filter(x=>x!==v);page=1;render();}));
}
for(const id of Object.keys(selections)){$(id+'Search').addEventListener('input',()=>drawChoices(id));drawChoices(id);}
function readFilters(){return {...Object.fromEntries(fields.map(k=>[k,$(k).value])),...Object.fromEntries(Object.entries(selections).map(([k,v])=>[k,[...v]]))};}
let page=1;
function resetFilters(){for(const k of fields)$(k).value='';for(const k of Object.keys(selections)){selections[k]=[];$(k+'Search').value='';drawChoices(k);}page=1;render();}
function group(rows,index,value){const m=new Map();for(const r of rows){const k=r[index];if(!m.has(k))m.set(k,[]);m.get(k).push(value(r));}return [...m].map(([name,vals])=>({name,vals}));}
function svg(content,w=700,h=340,label='Biểu đồ'){return `<svg viewBox="0 0 ${w} ${h}" role="img" aria-label="${esc(label)}"><title>${esc(label)}</title>${content}</svg>`;}
function empty(){return '<div class="empty">Không có quan sát phù hợp với bộ lọc.</div>';}
function bars(data,{unit='',percent=false}={}){
 if(!data.length)return empty();const w=700,h=Math.max(250,data.length*40+48),left=190,max=Math.max(...data.map(d=>d.value),1),right=610,scale=(right-left)/max;
 let s='';for(let i=0;i<=4;i++){const x=left+(right-left)*i/4;s+=`<line x1="${x}" y1="12" x2="${x}" y2="${h-28}" stroke="#e9e5df"/><text x="${x}" y="${h-8}" text-anchor="middle" class="tick">${percent?pct.format(max*i/4):vi.format(max*i/4)}</text>`;}
 data.forEach((d,i)=>{const y=18+i*40;s+=`<text x="${left-10}" y="${y+18}" text-anchor="end" class="label">${esc(d.name)}</text><rect data-n="${d.n}" x="${left}" y="${y}" height="26" width="${Math.max(0,d.value*scale)}" rx="3" fill="#537c80"><title>${esc(d.name)}: ${d.value} ${unit}; n=${d.n}</title></rect><text x="${Math.min(right+5,left+d.value*scale+7)}" y="${y+18}" class="label">${percent?pct.format(d.value):vi.format(d.value)}</text>`;});return svg(s,w,h,unit);
}
function histogram(prices){
 if(!prices.length)return empty();const bounds=[100000,200000,300000,400000,500000,750000,1000000,Infinity],labels=['≤100k','100–200k','200–300k','300–400k','400–500k','500–750k','750k–1tr','>1tr'],counts=Array(8).fill(0);
 prices.forEach(p=>counts[bounds.findIndex(b=>p<=b)]++);const max=Math.max(...counts,1);let s='';
 const ticks=Math.min(4,max);for(let i=0;i<=ticks;i++){const y=280-i*240/ticks;s+=`<line x1="42" y1="${y}" x2="686" y2="${y}" stroke="#e9e5df"/><text x="34" y="${y+4}" text-anchor="end" class="tick">${vi.format(max*i/ticks)}</text>`;}
 counts.forEach((n,i)=>{const x=55+i*79,h=n/max*240;s+=`<rect data-n="${n}" x="${x}" y="${280-h}" width="49" height="${h}" rx="3" fill="#cf7b5d"><title>${labels[i]}: ${n} snapshot</title></rect><text x="${x+24}" y="${272-h}" text-anchor="middle" class="label">${vi.format(n)}</text><text x="${x+24}" y="307" text-anchor="middle" class="tick">${labels[i]}</text>`;});return svg(s,700,330,'Phân bố giá hiện tại, VND');
}
function timeChart(rows){
 const dates=[...new Set(rows.map(r=>r[7]).filter(Boolean))].sort(),retailers=[...new Set(rows.map(r=>r[2]))].sort();if(!dates.length)return empty();
 const values=[];for(const n of retailers)for(const d of dates){const a=rows.filter(r=>r[2]===n&&r[7]===d).map(r=>r[5]);if(a.length)values.push(average(a));}const max=Math.max(...values,1);let s='';const x=i=>dates.length===1?365:55+i*610/(dates.length-1),y=v=>250-v/max*215;
 for(let i=0;i<=4;i++){const yy=y(max*i/4);s+=`<line x1="50" y1="${yy}" x2="675" y2="${yy}" stroke="#e9e5df"/><text x="44" y="${yy+4}" text-anchor="end" class="tick">${vi.format(max*i/4)}</text>`;}
 dates.forEach((d,i)=>s+=`<text x="${x(i)}" y="272" text-anchor="middle" class="tick">${d.slice(5)}</text>`);
 retailers.forEach((n,j)=>{let segment=[];const finish=()=>{if(segment.length>1)s+=`<polyline points="${segment.join(' ')}" fill="none" stroke="${colors[j%colors.length]}" stroke-width="2.5"/>`;segment=[];};dates.forEach((d,i)=>{const a=rows.filter(r=>r[2]===n&&r[7]===d).map(r=>r[5]);if(!a.length){finish();return;}const v=average(a);segment.push(`${x(i)},${y(v)}`);s+=`<circle cx="${x(i)}" cy="${y(v)}" r="4" fill="${colors[j%colors.length]}"><title>${esc(n)} ${d}: ${vi.format(v)} VND; n=${a.length}</title></circle>`;});finish();const lx=50+(j%3)*220,ly=299+Math.floor(j/3)*20;s+=`<circle cx="${lx}" cy="${ly}" r="4" fill="${colors[j%colors.length]}"/><text x="${lx+10}" y="${ly+4}" class="tick">${esc(n)}</text>`;});return svg(s,700,345,'Giá trung bình theo ngày và retailer, VND');
}
function qualityChart(rows){
 const groups=group(rows,2,r=>r);if(!groups.length)return empty();let s='';groups.forEach(({name,vals},i)=>{const y=20+i*48,n=vals.length,valid=vals.filter(r=>r[5]>0).length,listed=vals.filter(r=>r[5]>0&&r[6]>0).length;s+=`<text x="180" y="${y+15}" text-anchor="end" class="label">${esc(name)}</text><rect x="190" y="${y}" width="400" height="15" fill="#ede6df"/><rect data-n="${n}" x="190" y="${y}" width="${valid/n*400}" height="15" fill="#537c80"/><rect x="190" y="${y+19}" width="400" height="8" fill="#ede6df"/><rect x="190" y="${y+19}" width="${listed/n*400}" height="8" fill="#cf7b5d"/><text x="600" y="${y+14}" class="tick">${pct.format(valid/n*100)}%</text><text x="600" y="${y+30}" class="tick">${pct.format(listed/n*100)}%</text>`;});return svg(s,700,groups.length*48+40,'Coverage giá hiện tại (xanh) và cặp giá (cam), %');
}
function donut(rows){
 const parts=group(rows,3,r=>1).map(d=>({name:d.name,value:d.vals.length})).sort((a,b)=>b.value-a.value);if(!parts.length)return empty();let a=-Math.PI/2,s='';const cx=160,cy=160,r=110;
 parts.forEach((d,i)=>{const span=d.value/rows.length*Math.PI*2,b=a+span,large=span>Math.PI?1:0,x1=cx+r*Math.cos(a),y1=cy+r*Math.sin(a),x2=cx+r*Math.cos(b),y2=cy+r*Math.sin(b);if(parts.length===1)s+=`<circle cx="${cx}" cy="${cy}" r="${r}" fill="${colors[i]}"/>`;else s+=`<path d="M ${cx},${cy} L ${x1},${y1} A ${r},${r} 0 ${large},1 ${x2},${y2} Z" fill="${colors[i%colors.length]}"><title>${esc(d.name)}: ${d.value}</title></path>`;a=b;const y=40+i*37;s+=`<rect x="320" y="${y}" width="12" height="12" fill="${colors[i%colors.length]}"/><text x="342" y="${y+11}" class="label">${esc(d.name)} · ${pct.format(d.value/rows.length*100)}%</text>`;});s+=`<circle cx="${cx}" cy="${cy}" r="70" fill="#fffdf9"/><text x="${cx}" y="${cy}" text-anchor="middle" class="donut-total">${vi.format(rows.length)}</text><text x="${cx}" y="${cy+24}" text-anchor="middle" class="tick">snapshot</text>`;return svg(s,700,325,'Cơ cấu snapshot theo primary_category');
}
let lastState;
function render(){
 const f=readFilters(),state=calculate(DATA,f);lastState=state;const n=state.selected.length;
 $('scope').textContent=`${vi.format(n)} / ${vi.format(DATA.rows.length)} snapshot · ${new Set(state.selected.map(r=>r[2])).size} retailer · ${DATA.period}`;
 $('cards').dataset.matchingN=state.matching_n;$('cards').dataset.matchingOk=state.matching_ok;
 for(const id of Object.keys(selections))$(id+'Summary').textContent=selections[id].length?selections[id].length+' đã chọn':'Tất cả';
 const selectedLabels=Object.entries(selections).filter(([k,v])=>v.length).map(([k,v])=>({retailer:'Nhà bán lẻ',category:'Danh mục',brand:'Thương hiệu'}[k])+': '+v.map(x=>options[k].find(o=>o.value===x)?.label??x).join(', '));
 $('selectionSummary').textContent=selectedLabels.length?selectedLabels.join(' | '):'Chưa giới hạn nhà bán lẻ, danh mục hoặc thương hiệu. Có thể chọn nhiều giá trị trong mỗi nhóm.';
 const inverted=(f.min!==''&&f.max!==''&&Number(f.min)>Number(f.max))||(f.from&&f.to&&f.from>f.to);
 $('notice').textContent=inverted?'Khoảng lọc đảo ngược. Hãy nhập cận dưới không lớn hơn cận trên.':!n?'Không có dữ liệu phù hợp. Các KPI giá và tỷ lệ hiển thị trạng thái thiếu dữ liệu.':'';
 $('notice').hidden=!$('notice').textContent;
 $('cards').innerHTML=DATA.metricKeys.map((k,i)=>`<article class="card"><div class="card-label">${esc(DATA.metricNames[i])}</div><div class="card-value" data-metric="${k}" data-value="${state.metrics[k]===null?'null':state.metrics[k]}">${fmt(k,state.metrics[k])}</div><div class="card-unit">${['min','max','median','mean','gap'].includes(k)?'VND':k==='matching'?`${state.matching_ok}/${state.matching_n} trường hợp trong hàng đợi`:k==='discount'?`${state.checks.discountPairs} cặp giá hợp lệ`:k==='products'?'product_id duy nhất':k==='snapshots'?'dòng quan sát':`${state.checks.hist}/${n} giá hợp lệ`}</div></article>`).join('');
 for(const k of ['gap','matching']){const el=document.querySelector('[data-metric="'+k+'"]').closest('.card').querySelector('.card-unit');el.textContent=state.metrics[k]===null?(k==='gap'?DATA.gapReason:DATA.matchingReason):k==='gap'?'VND; mean of verified equivalent product/date gaps':'%';}
 $('cards').dataset.discountPairs=state.checks.discountPairs;
 $('hist').innerHTML=histogram(state.valid.map(r=>r[5]));$('hist').dataset.count=state.checks.hist;
 $('categoryChart').innerHTML=bars(group(state.valid,3,r=>r[5]).map(g=>({name:g.name,value:median(g.vals),n:g.vals.length})).sort((a,b)=>b.value-a.value),{unit:'VND'});$('categoryChart').dataset.count=state.checks.category;
 $('retailerChart').innerHTML=bars(group(state.valid,2,r=>r[5]).map(g=>({name:g.name,value:median(g.vals),n:g.vals.length})).sort((a,b)=>b.value-a.value),{unit:'VND'});
 $('timeChart').innerHTML=timeChart(state.valid);$('qualityChart').innerHTML=qualityChart(state.selected);$('qualityChart').dataset.count=state.checks.quality;
 const products=new Map();for(const r of state.selected){if(!products.has(r[0]))products.set(r[0],{r,prices:[],categories:new Set(),n:0,date:''});const p=products.get(r[0]);if(r[5]>0)p.prices.push(r[5]);p.categories.add(r[3]);p.n++;if(r[7]>p.date)p.date=r[7];}
 const list=[...products.values()].sort((a,b)=>a.r[8].localeCompare(b.r[8],'vi')),pages=Math.max(1,Math.ceil(list.length/25));page=Math.max(1,Math.min(page,pages));
 $('productCount').textContent=`${vi.format(list.length)} sản phẩm phù hợp · Trang ${page}/${pages} · 25 sản phẩm/trang`;
 $('dataPreview').innerHTML=list.slice((page-1)*25,page*25).map(p=>{const low=p.prices.length?Math.min(...p.prices):null,high=p.prices.length?Math.max(...p.prices):null;return `<tr data-product="${esc(p.r[0])}"><td><strong>${esc(p.r[8])}</strong><div class="source">${esc(p.r[0])}</div></td><td>${esc(p.r[2])}</td><td>${esc([...p.categories].join('; '))}</td><td>${esc(p.r[4])}</td><td>${low===null?'Thiếu':vi.format(low)+(low===high?'':' – '+vi.format(high))}</td><td>${p.n}</td><td>${esc(p.date)}</td></tr>`;}).join('');
 $('previous').disabled=page===1;$('next').disabled=page===pages;
}
for(const k of fields)for(const event of ['input','change'])$(k).addEventListener(event,()=>{page=1;render();});
$('reset').addEventListener('click',resetFilters);
$('previous').addEventListener('click',()=>{page--;render();});
$('next').addEventListener('click',()=>{page++;render();});
render();
"""

PAGE = r"""<!doctype html>
<html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>BEAUTYPRICE INTELLIGENCE — Dashboard EDA</title>
<style>
:root{--ink:#243b3a;--muted:#63716d;--paper:#f6f3ed;--panel:#fffdf9;--line:#deded4;--accent:#bc674b;--sage:#537c80}*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:15px/1.5 'Segoe UI',Arial,sans-serif}header{padding:38px 4vw 28px;background:#e9eee6;border-bottom:1px solid var(--line)}.eyebrow{font-size:12px;letter-spacing:.18em;text-transform:uppercase;color:var(--sage);font-weight:700}h1{font-family:Georgia,serif;font-weight:500;font-size:clamp(30px,3.3vw,49px);line-height:1.15;margin:11px 0}header p{max-width:940px;color:var(--muted);margin:12px 0 0}nav{display:flex;gap:23px;margin-top:22px}a{color:var(--ink);text-underline-offset:5px}main{max-width:1600px;margin:auto;padding:24px 4vw 60px}.filter-panel{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:20px;box-shadow:0 5px 25px #203b2810}.filter-grid{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:13px}label{display:block;font-size:12px;color:var(--muted);font-weight:600}select,input{display:block;width:100%;margin-top:5px;border:1px solid #cfd5cc;border-radius:6px;background:white;padding:9px 10px;color:var(--ink);font:inherit;font-size:13px;min-height:39px}select:focus,input:focus,button:focus-visible{outline:2px solid var(--sage);outline-offset:2px}.filter-actions{display:flex;align-items:center;gap:15px;margin-top:17px;flex-wrap:wrap}button{font:inherit;font-size:13px;border:1px solid var(--sage);border-radius:6px;padding:9px 15px;background:var(--sage);color:white;cursor:pointer}button.secondary{background:transparent;color:var(--ink)}.scope{font-size:12px;color:var(--muted)}.notice{margin:16px 0;padding:14px;border-radius:7px;background:#fff0d5;color:#755023}.cards{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:13px;margin:22px 0}.card{background:var(--panel);border:1px solid var(--line);border-radius:9px;padding:17px;min-height:132px}.card-label{font-size:12px;color:var(--muted);min-height:36px}.card-value{font-size:clamp(20px,2vw,30px);font-weight:600;line-height:1.2;letter-spacing:-.02em;margin:7px 0}.card-unit{font-size:11px;color:var(--muted)}.section-head{display:flex;justify-content:space-between;align-items:baseline;gap:20px;margin:35px 0 15px}h2{font-family:Georgia,serif;font-size:28px;font-weight:500;margin:0}.section-head p{color:var(--muted);font-size:12px;margin:0}.charts{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:18px}.chart-panel{background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:20px;min-width:0}.chart-panel h3{font-weight:600;font-size:16px;margin:0}.chart-panel p{font-size:12px;color:var(--muted);margin:5px 0 12px;min-height:36px}.chart svg{display:block;width:100%;height:auto}.chart{min-height:180px}.tick{font:11px 'Segoe UI',Arial,sans-serif;fill:#61706c}.label{font:12px 'Segoe UI',Arial,sans-serif;fill:#243b3a}.donut-total{font:26px Georgia,serif;fill:#243b3a}.wide{grid-column:1/-1}.wide .chart{max-width:850px}.empty{padding:55px 10px;text-align:center;color:var(--muted);background:#f4f4ef;border-radius:6px}.note{font-size:12px;color:var(--muted);padding:15px 0;max-width:1100px}.insights{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:16px}.insight{padding:20px;border-top:3px solid #c6a75d;background:var(--panel);border-radius:0 0 8px 8px}.insight h3{font-size:15px;margin:0 0 12px}.insight p{font-size:13px;margin:8px 0}.source{font-size:11px!important;color:var(--muted)}.table-wrap{overflow:auto;background:var(--panel);border:1px solid var(--line);border-radius:8px}table{border-collapse:collapse;width:100%;font-size:12px}th{text-align:left;background:#e9eee6;font-weight:600;padding:11px}td{padding:10px 11px;border-bottom:1px solid #e9e7e0;vertical-align:top}.pass{color:#37694d;font-weight:700}.fail{color:#a72f23;font-weight:700}details{margin-top:20px}summary{cursor:pointer;font-weight:600}.test-head{display:flex;gap:20px;align-items:center;flex-wrap:wrap;margin-bottom:15px}.test-head p{font-size:13px;color:var(--muted)}footer{margin-top:35px;border-top:1px solid var(--line);padding-top:20px;color:var(--muted);font-size:12px}@media(max-width:1000px){.filter-grid{grid-template-columns:repeat(3,minmax(0,1fr))}.cards{grid-template-columns:repeat(3,minmax(0,1fr))}}@media(max-width:680px){header{padding:27px 18px}main{padding:18px}.filter-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.cards{grid-template-columns:repeat(2,minmax(0,1fr))}.charts,.insights{grid-template-columns:1fr}.wide{grid-column:auto}.section-head{display:block}.section-head p{margin-top:7px}nav{gap:16px}.card{padding:13px}.card-value{font-size:22px}}
.multi{position:relative;margin:0;border:1px solid #cfd5cc;border-radius:6px;background:white;min-height:60px}.multi summary{padding:9px 10px;font-size:12px;display:flex;justify-content:space-between;gap:8px}.multi summary span{color:var(--sage);font-size:11px}.multi .dropdown{position:absolute;z-index:10;top:100%;left:0;background:white;border:1px solid #cfd5cc;box-shadow:0 8px 22px #203b2825;border-radius:6px;padding:10px;width:max(100%,300px);max-width:85vw}.choices{max-height:235px;overflow:auto}.choice{display:flex;align-items:flex-start;gap:8px;padding:8px 4px;font-size:12px;font-weight:400;color:var(--ink);cursor:pointer}.choice input{display:inline-block;width:16px;height:16px;min-height:16px;flex-shrink:0;margin:1px 0 0;padding:0}.choice:hover{background:#eef2eb}.choice-count{font-size:11px;color:var(--muted);margin-top:7px}.selections{font-size:12px;color:var(--muted);margin:12px 0 0;overflow-wrap:anywhere}.products-title{margin:25px 0 12px}.pagination{display:flex;align-items:center;gap:15px;margin:12px 0}button:disabled{opacity:.4;cursor:default}#dataPreview strong{font-weight:600}#dataPreview td:first-child{min-width:260px}#dataPreview td{max-width:400px}
</style></head><body>
<header><div class="eyebrow">BEAUTYPRICE INTELLIGENCE / GIAI ĐOẠN 6</div><h1>Hiểu giá từ dữ liệu quan sát.</h1><p>Khám phá mặt bằng giá, khuyến mãi và độ đầy đủ dữ liệu mỹ phẩm từ năm nhà bán lẻ tại Việt Nam. Mỗi KPI và biểu đồ sử dụng cùng bộ lọc.</p><nav><a href="#analysis">Biểu đồ</a><a href="#products">Sản phẩm</a></nav></header>
<main><section class="filter-panel" aria-label="Bộ lọc phân tích"><div class="filter-grid">
__MULTI__
<label>Giá từ (VND, gồm biên)<input id="min" type="number" min="0" placeholder="Không giới hạn" aria-label="Giá từ"></label>
<label>Giá đến (VND, gồm biên)<input id="max" type="number" min="0" placeholder="Không giới hạn" aria-label="Giá đến"></label>
<label>Từ ngày (UTC+7)<input id="from" type="date" aria-label="Từ ngày"></label>
<label>Đến ngày (UTC+7)<input id="to" type="date" aria-label="Đến ngày"></label>
<label>Chất lượng giá<select id="quality" aria-label="Chất lượng giá"><option value="">Tất cả trạng thái giá</option><option value="valid">Giá hợp lệ &gt; 0</option><option value="invalid">Giá thiếu / không hợp lệ</option></select></label>
<label>Trạng thái đối chiếu<select id="match" aria-label="Trạng thái đối chiếu"><option value="">Tất cả sản phẩm</option><option value="pending">Trong hàng đợi chưa xác minh</option><option value="verified">Đã matched và xác minh</option></select></label>
</div><div class="filter-actions"><button id="reset" class="secondary">Xóa bộ lọc</button><span id="scope" class="scope"></span></div><p id="selectionSummary" class="selections"></p><p class="source">Trong mỗi nhóm: chọn một hoặc nhiều giá trị để lấy hợp (OR). Giữa các nhóm: lấy giao (AND). Không chọn gì trong một nhóm nghĩa là tất cả. Nếu giữ sản phẩm rồi đổi nguồn không có sản phẩm đó, kết quả có thể rỗng.</p></section>
<div id="notice" class="notice" role="status" aria-live="polite" hidden></div><section id="cards" class="cards" aria-label="10 KPI"></section>
<p class="note">Giá trung bình, trung vị và chênh lệch giá mô tả toàn catalog trong phạm vi lọc. Tỷ lệ đối chiếu đo hàng đợi hiện có theo product_id + retailer_id; khi lọc, chỉ giữ trường hợp có quan sát phù hợp. Khi mẫu số bằng 0, tỷ lệ để trạng thái “Không có dữ liệu”.</p>
<section id="analysis"><div class="section-head"><h2>Giá, khuyến mãi và chất lượng</h2><p>Nguồn: price_snapshots (1).csv · Bộ lọc phía trên áp dụng cho tất cả biểu đồ</p></div><div class="charts">
<article class="chart-panel"><h3>Giá tập trung ở khoảng nào?</h3><p>Số snapshot theo khoảng giá (VND). Các khoảng sau gồm cận trên và loại cận dưới; chỉ giá hợp lệ.</p><div id="hist" class="chart"></div></article>
<article class="chart-panel"><h3>Danh mục nào có trung vị giá cao hơn?</h3><p>Trung vị giá snapshot, VND; giữ cả nhóm chưa phân loại. Chưa chuẩn hóa dung tích hoặc phân khúc.</p><div id="categoryChart" class="chart"></div></article>
<article class="chart-panel"><h3>Mặt bằng catalog theo nhà bán lẻ</h3><p>Trung vị giá snapshot, VND. Cơ cấu sản phẩm khác nhau; chưa thể kết luận nơi bán rẻ nhất.</p><div id="retailerChart" class="chart"></div></article>
<article class="chart-panel"><h3>Giá trung bình trong đợt thu thập</h3><p>VND theo ngày và retailer; chấm là điểm quan sát. Khác cơ cấu thu thập có thể tạo khác biệt giữa ngày.</p><div id="timeChart" class="chart"></div></article>
<article class="chart-panel"><h3>Nguồn nào đủ dữ liệu giá?</h3><p>Tỷ lệ giá hiện tại hợp lệ (xanh) và đủ cặp giá (cam), %. Mẫu số là toàn bộ snapshot của từng retailer trong bộ lọc.</p><div id="qualityChart" class="chart"></div></article>
</div></section>
<section id="products"><div class="section-head products-title"><h2>Danh sách sản phẩm theo bộ lọc</h2><p id="productCount"></p></div><div class="table-wrap"><table><thead><tr><th>Tên sản phẩm / mã</th><th>Nhà bán lẻ</th><th>Danh mục chính</th><th>Thương hiệu</th><th>Giá min–max trong bộ lọc (VND)</th><th>Số snapshot</th><th>Ngày cuối trong bộ lọc</th></tr></thead><tbody id="dataPreview"></tbody></table></div><div class="pagination"><button id="previous" class="secondary">Trang trước</button><button id="next" class="secondary">Trang sau</button></div></section>
<footer>Dashboard chạy cục bộ. Nguồn: products (1).csv, price_snapshots (1).csv, product_matching.csv. Mỗi dòng bảng sản phẩm tổng hợp các snapshot của một product_id trong bộ lọc; giá thống kê vẫn tính theo snapshot.</footer>
</main><script id="dataset" type="application/json">__DATA__</script><script>const DATA=JSON.parse(document.getElementById('dataset').textContent);
__ENGINE__
__UI__
</script></body></html>"""

def find_node():
    node = shutil.which('node')
    if node:
        return node
    bundled = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    return str(bundled) if bundled.is_file() else None

# Extend the working page with an independent review tab for both sources.
PAGE=PAGE.replace('<main>','<main><div role="tablist" aria-label="Phân tích"><button id="priceTab" role="tab">PRICE</button> <button id="reviewsTab" role="tab">REVIEWS</button></div><div id="priceView">',1)
PAGE=PAGE.replace('<footer>','</div>'+review_panel()+'<footer>',1)
PAGE=PAGE.replace('Nếu giữ sản phẩm rồi đổi nguồn không có sản phẩm đó, kết quả có thể rỗng.','')
PAGE=PAGE.replace('<section id="analysis">','<p class="note" id="discountDefinition">Discount KPI: trung bình discount_percent_recalculated = (giá niêm yết − giá hiện tại) / giá niêm yết × 100%. Chỉ cặp giá dương và giá niêm yết ≥ giá hiện tại. discount_percent trong CSV là giá trị đã xử lý; giá website gốc được giữ ở discount_percent_source; discount_percent_calculated là trường tính sẵn từ bước trước.</p><section id="analysis">',1)
PAGE=PAGE.replace('Giá trung bình, trung vị và chênh lệch giá mô tả toàn catalog trong phạm vi lọc. Tỷ lệ đối chiếu đo hàng đợi hiện có theo product_id + retailer_id; khi lọc, chỉ giữ trường hợp có quan sát phù hợp. Khi mẫu số bằng 0, tỷ lệ để trạng thái “Không có dữ liệu”.','Giá trung bình và trung vị mô tả snapshot/catalog trong phạm vi lọc. Cross-source Price Gap chỉ dùng sản phẩm tương đương đã xác minh ở ít nhất hai retailer cùng ngày. Successful Matching Rate chưa tính được khi chưa định nghĩa chính thức N_need_matching.')
PAGE=PAGE.replace('</style>','h1,h2,h3{font-family:"Segoe UI",Arial,sans-serif;letter-spacing:normal;font-feature-settings:normal;font-variant:normal;text-transform:none}[hidden]{display:none!important}[role=tablist]{margin-bottom:20px}[role=tab][aria-selected=true]{background:#243b3a}.card-value[data-value="null"]{font-size:18px;overflow-wrap:anywhere}.card-unit{overflow-wrap:anywhere}</style>',1)
ENGINE+=REVIEW_ENGINE
UI+=REVIEW_UI
add_review_charts(wb,reviews)

parser = argparse.ArgumentParser(description='Tạo biểu đồ, insight và dashboard BEAUTYPRICE; đối chiếu KPI với Python độc lập.')
parser.add_argument('--serve-qa', action='store_true', help='Mở máy chủ localhost có trang kiểm thử RIÊNG ở /__qa__; Ctrl+C để dừng.')
parser.add_argument('--port', type=int, default=8769)
args = parser.parse_args()

node = find_node()
review_tests,review_actuals,review_test_rows=review_engine_tests(payload,node,reviews)
engine_results = {}
if node:
    runner = ENGINE + "\nconst fs=require('fs');const d=JSON.parse(fs.readFileSync(0,'utf8'));console.log(JSON.stringify(d.tests.map(t=>({id:t.id,actual:calculate(d,t.filters)}))));"
    process = subprocess.run([node, '-e', runner], input=json.dumps({**payload, 'tests': test_cases}, ensure_ascii=False), text=True, encoding='utf-8', capture_output=True, timeout=90)
    if process.returncode:
        raise RuntimeError('Engine JavaScript không chạy được: ' + process.stderr[:1000])
    engine_results = {r['id']: r['actual'] for r in json.loads(process.stdout)}

test_summary, test_details = [], []
for test in test_cases:
    actual = engine_results.get(test['id'])
    expected = test['expected']
    statuses = []
    for k, name in zip(metric_keys, metric_names):
        e = expected['metrics'][k]
        a = actual['metrics'][k] if actual else None
        passed = actual is not None and ((a is None and e is None) or (a is not None and e is not None and math.isclose(a, e, rel_tol=1e-10, abs_tol=1e-8)))
        statuses.append(passed)
        test_details.append([test['id'], name, json.dumps(test['filters'], ensure_ascii=False), e, a, (a-e) if a is not None and e is not None else None, 'PASS' if passed else 'FAIL' if actual else 'NOT_RUN', 'Python csv/statistics độc lập so với chính engine JavaScript được nhúng trong dashboard'])
    charts_ok = bool(actual and actual['checks'] == expected['checks'] and actual['matching_n'] == expected['matching_n'] and actual['matching_ok'] == expected['matching_ok'])
    status = 'PASS' if all(statuses) and charts_ok else 'FAIL' if actual else 'NOT_RUN'
    test_summary.append([test['id'], test['name'], json.dumps(test['filters'], ensure_ascii=False), expected['metrics']['snapshots'], actual['metrics']['snapshots'] if actual else None, '10 KPI + 4 nhóm số đếm dùng cho 5 biểu đồ + tử/mẫu matching', status, 'Node thực thi engine dashboard; kết quả từng KPI ở Doi_chieu_KPI'])

if engine_results and any(row[6] != 'PASS' for row in test_summary):
    raise AssertionError('KPI dashboard không khớp Python độc lập; xem test case lỗi trước khi xuất kết quả.')

# Đối chiếu mặc định với bộ 10 KPI đã chốt ở nhiệm vụ 1–2.
baseline = load_workbook(HERE / '02_kpi_dictionary.xlsx', read_only=True, data_only=True)
baseline_kpis = {r[0]: r[7] for r in baseline['KPI_dictionary'].iter_rows(min_row=5, values_only=True)}
baseline_ids = ['KPI13', 'KPI14', 'KPI02', 'KPI03', 'KPI04', 'KPI08', 'KPI12', 'KPI05', 'KPI06', 'KPI09']
default_metrics = test_cases[0]['expected']['metrics']
for key_name, kpi_id in zip(metric_keys, baseline_ids):
    v, b = default_metrics[key_name], baseline_kpis[kpi_id]
    if v is None or b is None:
        assert v is b, (kpi_id, v, b)
    else:
        assert math.isclose(v, b, abs_tol=1e-8, rel_tol=1e-10), (kpi_id, v, b)
baseline.close()

report_path = HERE / '04_kiem_thu_dashboard.xlsx'

def same_number(actual, expected):
    return actual is None and expected is None or actual is not None and expected is not None and math.isclose(actual, expected, rel_tol=1e-10, abs_tol=1e-8)

def save_test_report(browser_cases=None,browser_reviews=None):
    browser_cases = browser_cases or []
    by_id = {r['id']: r for r in browser_cases}
    report = Workbook()
    report.remove(report.active)
    summary, details = [], []
    for test, original in zip(test_cases, test_summary):
        b = by_id.get(test['id'])
        b_metrics = b.get('actual', {}).get('metrics', {}) if b else {}
        same_metrics = bool(b and all(k in b_metrics and same_number(b_metrics[k], test['expected']['metrics'][k]) for k in metric_keys))
        same_checks = bool(b and b['actual']['checks'] == test['expected']['checks'] and b['actual']['matching_n'] == test['expected']['matching_n'] and b['actual']['matching_ok'] == test['expected']['matching_ok'])
        ui_ok = bool(same_metrics and same_checks and b.get('noDataOkay') and b.get('chartOkay') and b.get('separateOkay') and b.get('productTableOkay'))
        ui_status = 'PASS' if ui_ok else 'FAIL' if b else 'NOT_RUN'
        evidence = json.dumps(b, ensure_ascii=False, separators=(',', ':')) if b else 'Chưa chạy trên giao diện. Chạy --serve-qa, mở /__qa__ và thực thi bộ kiểm thử riêng.'
        summary.append(original + [ui_status, 'Giá trị DOM chi tiết và log ở các sheet sau' if b else evidence])
        for k, name in zip(metric_keys, metric_names):
            e = test['expected']['metrics'][k]
            a = engine_results.get(test['id'], {}).get('metrics', {}).get(k)
            d = b_metrics.get(k)
            details.append([test['id'], name, json.dumps(test['filters'], ensure_ascii=False), e, a, d, a-e if a is not None and e is not None else None, d-e if d is not None and e is not None else None, 'PASS' if engine_results and same_number(a,e) else 'FAIL' if engine_results else 'NOT_RUN', 'PASS' if b and k in b_metrics and same_number(d,e) else 'FAIL' if b else 'NOT_RUN'])
    report_table(report.create_sheet('Test_case'), ['Mã', 'Thao tác', 'Bộ lọc', 'Snapshot dự kiến', 'Snapshot engine', 'Phạm vi', 'Engine', 'Bằng chứng engine', 'Giao diện', 'Bằng chứng giao diện'], summary, [10, 48, 60, 20, 20, 48, 15, 56, 15, 64])
    detail_sheet = report.create_sheet('Doi_chieu_KPI')
    report_table(detail_sheet, ['Mã', 'KPI', 'Bộ lọc', 'Dự kiến Python', 'Engine JS', 'DOM dashboard', 'Sai lệch engine', 'Sai lệch DOM', 'Engine', 'DOM'], details, [10, 46, 60, 22, 22, 22, 22, 22, 15, 15])
    for i, row in enumerate(details, 2):
        fmt = '0.00%' if row[1].startswith('Tỷ lệ') else '#,##0' if row[1] in ('Tổng sản phẩm duy nhất', 'Tổng bản ghi giá') else '#,##0.00'
        for j in [4,5,6,7,8]: detail_sheet.cell(i,j).number_format = fmt
    log_rows = [[r['id'], json.dumps(r, ensure_ascii=False)] for r in browser_cases]
    if not log_rows: log_rows = [['NOT_RUN', 'Không tái sử dụng PASS của lần chạy cũ. Cần thực thi giao diện bản hiện tại.']]
    report_table(report.create_sheet('Log_giao_dien'), ['Mã', 'Log DOM/SVG, trạng thái không dữ liệu và bảng sản phẩm'], log_rows, [12, 165])
    methods = [
        ['Nguồn', ', '.join(payload['sourceHashes'])],
        ['Băm dữ liệu SHA256', json.dumps(payload['sourceHashes'], ensure_ascii=False)],
        ['Băm dashboard SHA256', hashlib.sha256(html_path.read_bytes()).hexdigest() if html_path.is_file() else 'Chưa xuất'],
        ['Đối chiếu độc lập', 'Python csv/statistics so với engine JavaScript; GUI đọc thẻ KPI DOM và số quan sát của các cột SVG thực sự được vẽ.'],
        ['Bộ lọc nhiều lựa chọn', 'OR trong retailer/category/brand; AND giữa các nhóm; danh sách rỗng nghĩa là tất cả. Dashboard không có bộ lọc sản phẩm.'],
        ['Dữ liệu rỗng', 'Số đếm bằng 0; thống kê và tỷ lệ thiếu mẫu số không được gán 0 giả.'],
        ['Kiểm tra thành phần', 'Dashboard không có insight hoặc khối kiểm thử; bảng sản phẩm có tên/mã và cập nhật theo bộ lọc.'],
        ['Kiểm tra 2 phần trước', '10 KPI mặc định khớp workbook KPI riêng; EDA và KPI được tách file.'],
        ['Sai số', 'abs_tol=1e-8; rel_tol=1e-10; VND hiển thị làm tròn đơn vị, tỷ lệ hiển thị 2 chữ số.'],
    ]
    report_table(report.create_sheet('Phuong_phap'), ['Mục', 'Nội dung'], methods, [36, 135])
    review_dom={r['id']:r for r in (browser_reviews or [])}
    review_rows=[]
    for row in review_test_rows:
        test_dom=review_dom.get(row[0]); actual_dom=test_dom.get('metrics',{}).get(row[3]) if test_dom else None
        ui_pass=bool(test_dom and row[3] in test_dom['metrics'] and same_number(actual_dom,row[4]) and test_dom.get('chartsOkay') and test_dom.get('noDataOkay'))
        review_rows.append(row[:7]+['PASS' if ui_pass else 'FAIL' if test_dom else 'NOT_RUN',actual_dom])
    report_table(report.create_sheet('Review_Test_Cases'),['Mã (REVIEW scope)','Thao tác','Bộ lọc review','Metric','Expected dataframe','Actual JS','Engine','UI','DOM value'],review_rows,[22,42,50,22,24,24,15,15,24])
    report_table(report.create_sheet('Review_UI_Log'),['Review test','DOM evidence'],[[r['id'],json.dumps(r,ensure_ascii=False)] for r in (browser_reviews or [])] or [['NOT_RUN','No browser review test executed']],[22,140])
    report['Phuong_phap'].append(['Matching definition',MATCH_REASON])
    report['Phuong_phap'].append(['Cross-source gap',GAP_REASON if not verified_identity else 'Verified equivalent group/date only; latest snapshot per retailer'])
    report['Phuong_phap'].append(['Review scope','Hasaki & Guardian; review T09–T16 are scoped to REVIEW, separate from existing PRICE test IDs.'])
    report['Phuong_phap'].append(['Queue counts','matching_n/matching_ok are queue diagnostics only, never N_need_matching.'])
    report.save(report_path)
    return all(row[8] == 'PASS' for row in summary)

html_path = HERE / '02_dashboard.html'
encoded_data = json.dumps(payload, ensure_ascii=False, separators=(',', ':')).replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
multi_markup = ''.join(f'<details class="multi"><summary>{label}<span id="{key}Summary">Tất cả</span></summary><div class="dropdown"><label>Tìm {label.lower()}<input id="{key}Search" type="search" aria-label="Tìm {label.lower()}" placeholder="Gõ tên để tìm"></label><div id="{key}" class="choices" role="group" aria-label="Chọn nhiều {label.lower()}"></div><div id="{key}Count" class="choice-count"></div></div></details>' for key,label in [('retailer','Nhà bán lẻ'),('category','Danh mục chính'),('brand','Thương hiệu')])
html_path.write_text(PAGE.replace('__DATA__', encoded_data).replace('__ENGINE__', ENGINE).replace('__UI__', UI).replace('__MULTI__', multi_markup), encoding='utf-8')
insight_path.write_text(insight_text, encoding='utf-8')
wb.save(OUTPUT_FILE)
save_test_report()
print(f'Created: {OUTPUT_FILE}')
print(f'Created: {html_path}')
print(f'Created: {insight_path}')
print(f'Created: {report_path}')
print(f'5 biểu đồ giá + 2 biểu đồ review Hasaki & Guardian trong Excel/dashboard; {len(insight_rows)} insight; {len(test_summary)} price tests; {len(review_tests)} review tests.')
print('Review engine: '+(f'{len(review_tests)}/{len(review_tests)} PASS' if review_actuals else 'NOT_RUN'))
print('Engine: ' + (f'{len(test_cases)}/{len(test_cases)} PASS; khớp bộ 10 KPI đã chốt' if engine_results else 'NOT_RUN — thiếu Node.js'))

# Trang QA chỉ được phục vụ trong bộ nhớ khi yêu cầu --serve-qa.
# Không chèn kỳ vọng kiểm thử hoặc giao diện kiểm thử vào dashboard bàn giao.
QA_PAGE = r'''<!doctype html><html lang="vi"><meta charset="utf-8">
<title>Kiểm thử riêng — BEAUTYPRICE</title><style>
body{font:15px/1.5 Segoe UI,sans-serif;margin:24px;color:#243b3a}button{padding:12px;background:#537c80;color:white;border:0;border-radius:6px;cursor:pointer}iframe{width:100%;height:650px;border:1px solid #ddd;margin-top:15px}pre{white-space:pre-wrap}table{border-collapse:collapse}td,th{padding:5px 15px;border:1px solid #ddd}
</style><h1>Kiểm thử dashboard — trang riêng</h1>
<p>Thao tác checkbox/input trên dashboard thật; đọc lại thẻ KPI DOM và số quan sát SVG. Đối chiếu với Python độc lập, sau đó lưu báo cáo 04. Không phải thành phần dashboard bàn giao.</p>
<button id="run">Chạy toàn bộ kiểm thử giao diện</button><p id="status">Đang tải dashboard…</p><div id="results"></div>
<iframe id="dashboard" src="/02_dashboard.html" title="Dashboard cần kiểm thử"></iframe>
<script>
const TESTS=__TESTS__, REVIEW_TESTS=__REVIEW_TESTS__, HASH=__HASH__;
const frame=document.getElementById('dashboard'),status=document.getElementById('status');
frame.onload=()=>status.textContent='Sẵn sàng kiểm thử '+TESTS.length+' trường hợp.';
const pause=()=>new Promise(resolve=>requestAnimationFrame(()=>requestAnimationFrame(resolve)));
const near=(a,b)=>a===null||b===null?a===b:Math.abs(a-b)<=Math.max(1e-8,Math.abs(b)*1e-10);
document.getElementById('run').onclick=async()=>{
 const button=document.getElementById('run');button.disabled=true;const output=[];
 try{
 for(const test of TESTS){
  status.textContent='Đang chạy '+test.id+' — '+test.name;
  const d=frame.contentDocument,w=frame.contentWindow;
  d.getElementById('reset').click();
  if(test.id==='T06'){
   const seed=[...d.querySelectorAll('#retailer input[type=checkbox]')].find(x=>x.value==='Hasaki');
   if(seed)seed.click();d.getElementById('reset').click();
  }
  for(const key of ['retailer','category','brand']){
   let values=test.filters[key]||[];if(!Array.isArray(values))values=[values];
   for(const value of values){
    const search=d.getElementById(key+'Search');search.value=value;search.dispatchEvent(new w.Event('input',{bubbles:true}));
    const box=[...d.querySelectorAll('#'+key+' input[type=checkbox]')].find(x=>x.value===value);
    if(!box)throw Error(test.id+': Không tìm thấy '+key+'='+value);
    box.click();
   }
  }
  for(const key of ['min','max','from','to','quality','match'])if(test.filters[key]!==undefined){
   const input=d.getElementById(key);input.value=String(test.filters[key]);input.dispatchEvent(new w.Event('input',{bubbles:true}));input.dispatchEvent(new w.Event('change',{bubbles:true}));
  }
  await pause();
  const metrics=Object.fromEntries([...d.querySelectorAll('[data-metric]')].map(x=>[x.dataset.metric,JSON.parse(x.dataset.value)]));
  const sum=id=>[...d.querySelectorAll('#'+id+' [data-n]')].reduce((s,x)=>s+Number(x.dataset.n),0);
  const checks={hist:sum('hist'),category:sum('categoryChart'),discountPairs:Number(d.getElementById('cards').dataset.discountPairs),quality:sum('qualityChart')};
  const cards=d.getElementById('cards');
  const actual={metrics,checks,matching_n:Number(cards.dataset.matchingN),matching_ok:Number(cards.dataset.matchingOk)};
  const noDataOkay=test.expected.metrics.snapshots!==0||(!d.getElementById('notice').hidden&&d.querySelector('[data-metric=min]').textContent.includes('Không có dữ liệu'));
  const chartOkay=[['hist',test.expected.checks.hist],['categoryChart',test.expected.checks.category],['retailerChart',test.expected.checks.hist],['timeChart',test.expected.checks.hist],['qualityChart',test.expected.checks.quality]].every(([id,n])=>Boolean(d.querySelector('#'+id+(n?' svg':' .empty'))));
  const separateOkay=!d.querySelector('#insights,#findings,#validation,#test,#runTests');
  const productRows=[...d.querySelectorAll('#dataPreview tr[data-product]')];
  const productTableOkay=test.expected.metrics.products===0?productRows.length===0:productRows.length>0&&productRows.every(r=>r.querySelector('strong')&&r.dataset.product&&r.querySelector('strong').textContent.length>0);
  const item={id:test.id,actual,noDataOkay,chartOkay,separateOkay,productTableOkay};
  const pass=Object.keys(test.expected.metrics).every(k=>near(metrics[k],test.expected.metrics[k]))&&JSON.stringify(checks)===JSON.stringify(test.expected.checks)&&actual.matching_n===test.expected.matching_n&&actual.matching_ok===test.expected.matching_ok&&noDataOkay&&chartOkay&&separateOkay&&productTableOkay;
  item.pass=pass;output.push(item);
  document.getElementById('results').textContent=output.map(x=>x.id+': '+(x.pass?'PASS':'FAIL')).join(' | ');
 }
 const reviewOutput=[];
 for(const test of REVIEW_TESTS){
  const d=frame.contentDocument,w=frame.contentWindow;
  d.getElementById('reviewsTab').click();d.getElementById('reviewReset').click();
  const source=d.getElementById('reviewSource');source.value=test.filters.source||'';source.dispatchEvent(new w.Event('change',{bubbles:true}));
  const input=d.getElementById('reviewCategory');input.value=test.filters.category||'';input.dispatchEvent(new w.Event('input',{bubbles:true}));await pause();
  const metrics=Object.fromEntries([...d.querySelectorAll('[data-review-metric]')].map(x=>[x.dataset.reviewMetric,JSON.parse(x.dataset.value)]));
  const sum=id=>[...d.querySelectorAll('#'+id+' [data-n]')].reduce((s,x)=>s+Number(x.dataset.n),0);
  const chartsOkay=test.expected.total>0?sum('reviewCategoryChart')===test.expected.total&&Boolean(d.querySelector('#ratingChart svg')):Boolean(d.querySelector('#ratingChart .empty,#reviewCategoryChart .empty'));
  const noDataOkay=test.expected.total>0||!d.getElementById('reviewNotice').hidden;
  const pass=Object.keys(test.expected).every(k=>near(metrics[k],test.expected[k]))&&chartsOkay&&noDataOkay;
  reviewOutput.push({id:test.id,metrics,chartsOkay,noDataOkay,pass});
 }
 frame.contentDocument.getElementById('reviewReset').click();frame.contentDocument.getElementById('priceTab').click();frame.contentDocument.getElementById('reset').click();
 const response=await fetch('/__qa__/results',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({dashboard_sha256:HASH,cases:output,reviewCases:reviewOutput})});
 const saved=await response.json();if(!response.ok)throw Error(saved.error);
 status.textContent=output.filter(x=>x.pass).length+'/'+TESTS.length+' price PASS; '+reviewOutput.filter(x=>x.pass).length+'/'+REVIEW_TESTS.length+' review PASS. '+saved.message;
 }catch(error){status.textContent='Lỗi kiểm thử: '+error.message;}finally{button.disabled=false;}
};</script></html>'''

if args.serve_qa:
    dashboard_hash = hashlib.sha256(html_path.read_bytes()).hexdigest()
    qa_bytes = QA_PAGE.replace('__TESTS__', json.dumps(test_cases, ensure_ascii=False).replace('<', '\\u003c')).replace('__REVIEW_TESTS__',json.dumps(review_tests,ensure_ascii=False).replace('<','\\u003c')).replace('__HASH__', json.dumps(dashboard_hash)).encode('utf-8')

    class QAHandler(SimpleHTTPRequestHandler):
        def __init__(self, *handler_args, **kwargs):
            super().__init__(*handler_args, directory=str(HERE), **kwargs)

        def do_GET(self):
            if self.path.split('?')[0] == '/__qa__':
                self.send_response(200)
                self.send_header('Content-Type', 'text/html; charset=utf-8')
                self.send_header('Content-Length', str(len(qa_bytes)))
                self.end_headers()
                self.wfile.write(qa_bytes)
            else:
                super().do_GET()

        def do_POST(self):
            try:
                if self.path != '/__qa__/results':
                    raise ValueError('Đường dẫn không được hỗ trợ')
                length = int(self.headers.get('Content-Length', '0'))
                if not 0 < length < 2_000_000:
                    raise ValueError('Kích thước log không hợp lệ')
                evidence = json.loads(self.rfile.read(length))
                if evidence.get('dashboard_sha256') != dashboard_hash or hashlib.sha256(html_path.read_bytes()).hexdigest() != dashboard_hash:
                    raise ValueError('Dashboard đã thay đổi; phải kiểm thử lại bản hiện tại')
                cases = evidence.get('cases', [])
                if len(cases) != len(test_cases) or {r['id'] for r in cases} != {r['id'] for r in test_cases}:
                    raise ValueError('Log không đủ toàn bộ trường hợp kiểm thử')
                review_cases=evidence.get('reviewCases',[])
                if len(review_cases)!=len(review_tests) or {r['id'] for r in review_cases}!={r['id'] for r in review_tests}:
                    raise ValueError('Incomplete review UI evidence')
                passed = save_test_report(cases,review_cases)
                response = {'message': 'Đã lưu 04_kiem_thu_dashboard.xlsx.', 'passed': passed}
                code = 200
            except (ValueError, KeyError, TypeError) as error:
                response, code = {'error': str(error)}, 400
            body = json.dumps(response, ensure_ascii=False).encode('utf-8')
            self.send_response(code)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    print(f'Trang kiểm thử RIÊNG: http://127.0.0.1:{args.port}/__qa__', flush=True)
    server = ThreadingHTTPServer(('127.0.0.1', args.port), QAHandler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print('Đã dừng máy chủ kiểm thử.')
    finally:
        server.server_close()

