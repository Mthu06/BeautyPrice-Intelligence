"""Hasaki and Guardian review tab and tests; independent of price data."""
import json
import math
import subprocess
import pandas as pd
from openpyxl.chart import BarChart, Reference
from stage6_data import review_metrics

REVIEW_ENGINE = r'''
function calculateReviews(data,f={}){
 const rows=data.reviews.filter(r=>(!f.source||r.source===f.source)&&(!f.category||r.category===f.category));
 const ratings=rows.map(r=>r.rating).filter(v=>Number.isFinite(v)&&v>=1&&v<=5);
 const rate=(k)=>{const a=rows.filter(r=>r[k]!==null&&r[k]!=='');return a.length?a.filter(r=>['yes','true','1'].includes(r[k])).length/a.length:null;};
 const unknown=rows.filter(r=>String(r.matching??'').toLowerCase()==='unknown').length;
 return {metrics:{total:rows.length,products:new Set(rows.map(r=>r.product).filter(Boolean)).size,mean:average(ratings),median:median(ratings),purchased:rate('bought'),images:rate('images'),unknown,unknown_rate:rows.length?unknown/rows.length:null},rows};
}
'''

REVIEW_UI = r'''
const reviewKeys=['total','products','mean','median','purchased','images','unknown','unknown_rate'];
const reviewNames=['Total reviews','Unique reviewed products','Average rating','Median rating','Purchased review rate','Review with image rate','Source product-match flag: unknown rows','Source product-match flag: unknown rate'];
function renderReviews(){
 const state=calculateReviews(DATA,{source:$('reviewSource').value,category:$('reviewCategory').value});
 const sourceLabel=$('reviewSource').selectedOptions[0].textContent;
 $('reviewCards').innerHTML=reviewKeys.map((k,i)=>{const v=state.metrics[k],text=v===null?'Không có dữ liệu':['purchased','images','unknown_rate'].includes(k)?pct.format(v*100)+'%':['mean','median'].includes(k)?pct.format(v)+' / 5':vi.format(v);return `<article class="card"><div class="card-label">${reviewNames[i]}</div><div class="card-value" data-review-metric="${k}" data-value="${JSON.stringify(v)}">${text}</div><div class="card-unit">${esc(sourceLabel)}</div></article>`;}).join('');
 const counts=new Map();for(const r of state.rows)if(r.rating!==null)counts.set(r.rating,(counts.get(r.rating)||0)+1);
 $('ratingChart').innerHTML=bars([...counts].sort((a,b)=>a[0]-b[0]).map(([r,n])=>({name:String(r)+' sao',value:n,n})),{unit:'reviews'});
 const cats=new Map();for(const r of state.rows)cats.set(r.category,(cats.get(r.category)||0)+1);
 $('reviewCategoryChart').innerHTML=bars([...cats].sort((a,b)=>b[1]-a[1]).map(([name,n])=>({name,value:n,n})),{unit:'reviews'});
 $('reviewNotice').hidden=state.rows.length>0;
 $('reviewScope').textContent=vi.format(state.rows.length)+' / '+vi.format(DATA.reviews.length)+' review · '+sourceLabel+'; mẫu số tỷ lệ mua hàng/ảnh là số dòng có trạng thái được ghi nhận.';
 const sourceFlags=[];
 for(const source of [...new Set(state.rows.map(r=>r.source))].sort()){
  const rows=state.rows.filter(r=>r.source===source),flags=new Map();
  for(const r of rows){const flag=r.matching===null||String(r.matching).trim()===''?'MISSING':r.matching;flags.set(flag,(flags.get(flag)||0)+1);}
  sourceFlags.push(source+': '+[...flags].sort((a,b)=>a[0].localeCompare(b[0])).map(([flag,n])=>flag+'='+vi.format(n)).join(', '));
 }
 $('sourceFlagScope').textContent='Source product-match flag — '+(sourceFlags.join('; ')||'Không có dữ liệu')+'. This source field is NOT a verified link to price products.'+($('reviewSource').value==='hasaki'?' '+state.metrics.unknown+' rows marked unknown in the Hasaki source field.':'');
 $('reviewSampleScope').textContent='Toàn mẫu thu thập: '+[...new Set(DATA.reviews.map(r=>r.source))].sort().map(source=>{const n=DATA.reviews.filter(r=>r.source===source).length;return source+' '+vi.format(n)+'/'+vi.format(DATA.reviews.length)+' ('+pct.format(n/DATA.reviews.length*100)+'%)';}).join('; ')+'. Pooled metrics chịu ảnh hưởng mạnh bởi Hasaki, không phải so sánh cân bằng hai retailer. Xem per-source metrics; không dùng pooled average để kết luận mức hài lòng khách hàng.';
 const knownBought=state.rows.filter(r=>r.bought!==null&&r.bought!=='');
 $('purchasedScope').textContent=knownBought.length&&knownBought.every(r=>['yes','true','1'].includes(r.bought))?'100% bản ghi review trong bộ lọc có trạng thái mua hàng được đánh dấu đã mua; không suy rộng cho toàn bộ khách hàng.':'Tỷ lệ đã mua chỉ mô tả cờ is_bought của các bản ghi review trong bộ lọc.';
}
$('reviewCategoryOptions').innerHTML=[...new Set(DATA.reviews.map(r=>r.category))].sort((a,b)=>a.localeCompare(b,'vi')).map(c=>`<option value="${esc(c)}"></option>`).join('');
$('reviewCategory').addEventListener('change',renderReviews);
$('reviewCategory').addEventListener('input',renderReviews);
$('reviewSource').addEventListener('change',renderReviews);
$('reviewReset').addEventListener('click',()=>{$('reviewSource').value='';$('reviewCategory').value='';renderReviews();});
function showTab(name){$('priceView').hidden=name!=='price';$('reviewView').hidden=name!=='reviews';$('priceTab').setAttribute('aria-selected',String(name==='price'));$('reviewsTab').setAttribute('aria-selected',String(name==='reviews'));}
$('priceTab').addEventListener('click',()=>showTab('price'));
$('reviewsTab').addEventListener('click',()=>showTab('reviews'));
renderReviews();showTab('price');
'''

def review_panel():
    return '''<section id="reviewView" hidden aria-label="Hasaki and Guardian reviews">
<div class="section-head"><h2>REVIEWS — Hasaki &amp; Guardian</h2></div>
<p class="note">Supporting review metrics — tách biệt với 10 core price KPI. Nguồn: hasaki_reviews.csv và guardian_reviews.csv; mỗi dòng là một review, sản phẩm được đếm theo nguồn + source_product_id. Danh mục giữ nguyên nhãn nguồn.<br>Review product identifiers/flags are source-specific and have not been verified as equivalent to price-product identifiers. Review analysis is therefore independent of price analysis.</p>
<p class="note" id="reviewSampleScope"></p>
<p class="note" id="purchasedScope"></p>
<div class="filter-panel"><label>Nguồn review<select id="reviewSource"><option value="">Tất cả — Hasaki &amp; Guardian</option><option value="hasaki">Hasaki</option><option value="guardian">Guardian</option></select></label><label>Danh mục review<input id="reviewCategory" list="reviewCategoryOptions" placeholder="Tất cả danh mục; chọn hoặc nhập đúng tên"><datalist id="reviewCategoryOptions"></datalist></label><div class="filter-actions"><button id="reviewReset" class="secondary">Xóa bộ lọc review</button><span id="reviewScope" class="scope"></span></div></div>
<p id="reviewNotice" class="notice" hidden>Không có review phù hợp với bộ lọc.</p>
<div id="reviewCards" class="cards"></div><p class="note" id="sourceFlagScope"></p><div class="charts">
<article class="chart-panel"><h3>Phân bố điểm đánh giá</h3><p>Observed rating distribution in the collected sample; gồm điểm nửa sao, không suy rộng cho toàn bộ khách hàng. Đơn vị: review; phạm vi: nguồn và danh mục review đã lọc.</p><div id="ratingChart" class="chart"></div></article>
<article class="chart-panel"><h3>Số review theo danh mục</h3><p>Số dòng review theo category nguồn; đơn vị review. Nguồn và danh mục áp dụng theo bộ lọc review.</p><div id="reviewCategoryChart" class="chart"></div></article>
</div></section>'''

def add_review_charts(workbook, reviews):
    chart_sheet=workbook.create_sheet('Review_Charts')
    data_sheet=workbook.create_sheet('Review_Data')
    chart_sheet['A1']='Hasaki & Guardian reviews — phân tích độc lập với giá'
    chart_sheet.merge_cells('A1:R1')
    chart_sheet['A2']='Rating: 1–5. Sources: hasaki_reviews.csv + guardian_reviews.csv; all source rows.'
    chart_sheet.merge_cells('A2:R2')
    data_sheet.append(['rating','review_n'])
    for rating,count in reviews['rating_numeric'].dropna().value_counts().sort_index().items(): data_sheet.append([float(rating),int(count)])
    end=data_sheet.max_row
    data_sheet['D1']='category'; data_sheet['E1']='review_n'
    for i,(category,count) in enumerate(reviews['category_display'].value_counts().items(),2): data_sheet.cell(i,4,str(category));data_sheet.cell(i,5,int(count))
    data_sheet['G1']='source'; data_sheet['H1']='review_n'
    for i,(source,count) in enumerate(reviews['source'].value_counts().items(),2): data_sheet.cell(i,7,source);data_sheet.cell(i,8,int(count))
    for title,start_col,count,anchor in [('Phân bố điểm đánh giá — Hasaki & Guardian',1,end,'A4'),('Số review theo danh mục — Hasaki & Guardian',4,reviews['category_display'].nunique()+1,'A21')]:
        c=BarChart(); c.type='col'; c.title=title;c.x_axis.title='Rating (stars)' if start_col==1 else 'Category';c.y_axis.title='Review count';c.legend=None
        c.add_data(Reference(data_sheet,min_col=start_col+1,min_row=1,max_row=count),titles_from_data=True)
        c.set_categories(Reference(data_sheet,min_col=start_col,min_row=2,max_row=count));c.width=25;c.height=9;chart_sheet.add_chart(c,anchor)
    for col in range(1,19):
        from openpyxl.utils import get_column_letter
        chart_sheet.column_dimensions[get_column_letter(col)].width=9
    data_sheet.column_dimensions['D'].width=32
    for s in (chart_sheet,data_sheet): s.sheet_view.showGridLines=False

def review_test_cases(reviews):
    categories=sorted(reviews['category_display'].unique())
    test_specs=[('T09','Both review datasets load successfully',{}),('T10','Review category filter',{'category':categories[0]}),('T11','Review KPI reconciliation',{'category':categories[-1]}),('T12','Review no-data case',{'category':'__NO_SUCH_REVIEW_CATEGORY__'}),('T13','Hasaki source filter',{'source':'hasaki'}),('T14','Guardian source filter',{'source':'guardian'}),('T15','Hasaki source and category',{'source':'hasaki','category':sorted(reviews.loc[reviews.source.eq('hasaki'),'category_display'].unique())[0]}),('T16','Guardian source and category',{'source':'guardian','category':sorted(reviews.loc[reviews.source.eq('guardian'),'category_display'].unique())[0]})]
    result=[]
    for tid,name,filters in test_specs:
        selected=reviews
        if filters.get('source'): selected=selected.loc[selected.source.eq(filters['source'])]
        if filters.get('category'): selected=selected.loc[selected.category_display.eq(filters['category'])]
        # Independent expected values computed from the filtered dataframe, not JS output.
        expected=review_metrics(selected)
        result.append(dict(id=tid,name=name,filters=filters,expected=expected))
    return result

def review_engine_tests(payload,node,reviews):
    tests=review_test_cases(reviews)
    if not node: return tests,{},[]
    runner='function average(a){return a.length?a.reduce((s,x)=>s+x,0)/a.length:null;} function median(a){if(!a.length)return null;const b=[...a].sort((x,y)=>x-y),i=Math.floor(b.length/2);return b.length%2?b[i]:(b[i-1]+b[i])/2;}'+REVIEW_ENGINE+"\nconst fs=require('fs');const d=JSON.parse(fs.readFileSync(0,'utf8'));console.log(JSON.stringify(d.tests.map(t=>({id:t.id,actual:calculateReviews(d,t.filters).metrics}))));"
    process=subprocess.run([node,'-e',runner],input=json.dumps({**payload,'tests':tests},ensure_ascii=False),text=True,encoding='utf-8',capture_output=True,timeout=60)
    if process.returncode: raise RuntimeError(process.stderr)
    actuals={r['id']:r['actual'] for r in json.loads(process.stdout)}
    rows=[]
    for t in tests:
        for key,expected in t['expected'].items():
            actual=actuals[t['id']][key]
            same=actual is None and expected is None or actual is not None and expected is not None and math.isclose(actual,expected,abs_tol=1e-10,rel_tol=1e-10)
            rows.append([t['id'],t['name'],json.dumps(t['filters'],ensure_ascii=False),key,expected,actual,'PASS' if same else 'FAIL','NOT_RUN'])
    if any(r[6]!='PASS' for r in rows): raise AssertionError('Review engine mismatch')
    return tests,actuals,rows
