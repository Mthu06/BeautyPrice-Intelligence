"""Run scripts/notebook cells and reconcile delivered outputs with independent CSV calculations."""
from pathlib import Path
import ast
import contextlib
import hashlib
import io
import json
import math
import re
import subprocess
import sys
import urllib.request
import pandas as pd
from openpyxl import load_workbook

HERE=Path(__file__).resolve().parent
DATA=HERE.parent/'datar'
sys.path.insert(0,str(HERE))
if hasattr(sys.stdout,'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')

def hashes():
    return {str(p.relative_to(DATA)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(DATA.rglob('*.csv'))}

def execute_notebook(path):
    notebook=json.loads(path.read_text(encoding='utf-8'))
    execution=0
    namespace={'__name__':'__main__'}
    outputs=[]
    def display(value):
        text=value.to_string(index=False) if isinstance(value,pd.DataFrame) else repr(value)
        data={'text/plain':[text]}
        if isinstance(value,pd.DataFrame): data['text/html']=[value.to_html(index=False)]
        outputs.append({'output_type':'display_data','data':data,'metadata':{}})
    namespace['display']=display
    for index,cell in enumerate(notebook['cells']):
        cell.setdefault('id',f'cell-{index:03}')
        if cell['cell_type']!='code': continue
        execution+=1; outputs=[]; stream=io.StringIO()
        tree=ast.parse(''.join(cell['source']))
        with contextlib.redirect_stdout(stream):
            if tree.body and isinstance(tree.body[-1],ast.Expr):
                expr=tree.body.pop();exec(compile(tree,str(path),'exec'),namespace)
                value=eval(compile(ast.Expression(expr.value),str(path),'eval'),namespace)
                if value is not None: display(value)
            else: exec(compile(tree,str(path),'exec'),namespace)
        if stream.getvalue(): outputs.insert(0,{'output_type':'stream','name':'stdout','text':stream.getvalue().splitlines(keepends=True)})
        cell['execution_count']=execution;cell['outputs']=outputs
    notebook['metadata']['execution_validation']='Code cells executed sequentially with Python AST runner; Jupyter kernel unavailable in validation environment.'
    path.write_text(json.dumps(notebook,ensure_ascii=False,indent=1),encoding='utf-8')
    return execution

def close(a,b):
    return a is None and b is None or a is not None and b is not None and math.isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-8)

def check_startup():
    """HTTP startup evidence only. Never infer UI acceptance from an HTTP response."""
    result_path=HERE/'05_validation_results.json'
    result=json.loads(result_path.read_text(encoding='utf-8'))
    checks=[]
    try:
        for route in ['/02_dashboard.html','/__qa__']:
            url='http://127.0.0.1:8769'+route
            with urllib.request.urlopen(url,timeout=10) as response:
                body=response.read()
                item=dict(url=url,http_status=response.status,bytes=len(body),sha256=hashlib.sha256(body).hexdigest())
            if route=='/02_dashboard.html':
                item['matches_saved_html']=body==(HERE/'02_dashboard.html').read_bytes()
                if not item['matches_saved_html']: raise AssertionError('Served dashboard does not match current saved HTML')
            if item['http_status']!=200: raise AssertionError('HTTP startup check failed')
            checks.append(item)
        startup=dict(status='PASS',scope='HTTP server startup only; not browser rendering or UI acceptance',checks=checks)
    except Exception as error:
        startup=dict(status='FAIL',scope='HTTP startup check; may be blocked by environment',error=str(error),checks=checks)
    result['dashboard_startup']=startup
    result_path.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(startup,ensure_ascii=False,indent=2))
    if startup['status']!='PASS': raise SystemExit(1)

def main():
    before=hashes()
    previous_path=HERE/'05_validation_results.json'
    previous=json.loads(previous_path.read_text(encoding='utf-8')) if previous_path.exists() else None
    # Capture accepted EDA values before regeneration; sheet renames must not change statistics.
    prior=load_workbook(HERE/'01_ket_qua_eda.xlsx',read_only=True,data_only=True)
    prior_numeric_name='Numeric_Descriptive' if 'Numeric_Descriptive' in prior else '02_Numeric_Descriptive'
    def numeric_values(sheet):
        headings=[c.value for c in sheet[3]]
        return [{k:v for k,v in zip(headings,row) if k!='filter_conditions'} for row in sheet.iter_rows(min_row=4,values_only=True)]
    prior_numeric=numeric_values(prior[prior_numeric_name])
    quality_name='Chat_luong_du_lieu' if 'Chat_luong_du_lieu' in prior else '11_Collection_Status_Quality'
    outlier_name='Ngoai_le_gia' if 'Ngoai_le_gia' in prior else '12_Price_Outlier_Candidates'
    prior_quality=list(prior[quality_name].iter_rows(min_row=4,values_only=True))
    prior_outliers=list(prior[outlier_name].iter_rows(min_row=4,values_only=True))
    prior_missing=list(prior['Ty_le_trong'].iter_rows(min_row=4,values_only=True)) if 'Ty_le_trong' in prior else []
    prior_extra_iqr=[[r[0],r[1],r[2],r[15],r[16],r[17]] for r in prior['Thong_ke_mo_ta'].iter_rows(min_row=5,values_only=True)] if 'Thong_ke_mo_ta' in prior else list(prior['10_Extra_IQR_Diagnostics'].iter_rows(min_row=4,values_only=True))
    prior.close()
    prior_kpi=load_workbook(HERE/'02_kpi_dictionary.xlsx',read_only=True,data_only=True)
    prior_core_kpi=list(prior_kpi['KPI_dictionary'].iter_rows(min_row=5,values_only=True))
    prior_kpi.close()
    runs={}
    for file in ['01_eda.py','02_kpi.py','02_ve_bieu_do.py']:
        proc=subprocess.run([sys.executable,'-X','utf8',str(HERE/file)],cwd=HERE,text=True,encoding='utf-8',capture_output=True)
        print(proc.stdout)
        if proc.returncode: raise RuntimeError(proc.stderr)
        runs[file]='PASS'
    notebooks={p.name:execute_notebook(p) for p in [HERE/'01_eda_beautyprice.ipynb',HERE/'02_kpi_beautyprice.ipynb']}
    snapshots=pd.read_csv(next(DATA.glob('price_snapshots*.csv')),low_memory=False)
    products=pd.read_csv(next(DATA.glob('products*.csv')),low_memory=False)
    matching=pd.read_csv(DATA/'product_matching.csv',low_memory=False)
    review_sources={source:pd.read_csv(next(DATA.rglob(source+'_reviews.csv')),dtype=str) for source in ['hasaki','guardian']}
    reviews=pd.concat(review_sources.values(),ignore_index=True)
    p=pd.to_numeric(snapshots.current_price_vnd,errors='coerce').replace([math.inf,-math.inf],math.nan)
    valid=p.notna() & p.gt(0); clean=p[valid]
    q1,q3=clean.quantile(.25),clean.quantile(.75);iqr=q3-q1
    outlier=((clean<q1-1.5*iqr)|(clean>q3+1.5*iqr)).sum()
    dates=pd.to_datetime(snapshots.collected_at,errors='coerce',utc=True).dt.tz_convert('Asia/Ho_Chi_Minh')
    rating=pd.to_numeric(reviews.rating,errors='coerce');rating=rating[rating.between(1,5)]
    review_dates=pd.to_datetime(reviews.review_date.str.replace(r'^(\d{1,2})\s+tháng\s+(\d{1,2}),?\s+(\d{4})$',r'\3-\2-\1',regex=True),format='mixed',errors='coerce')
    unknown=reviews.product_id_match.fillna('').str.lower().eq('unknown')
    price_result=dict(total_snapshots=len(snapshots),unique_snapshot_products=snapshots.product_id.nunique(),master_products=products.product_id.nunique(),brands=snapshots.brand.nunique(),master_brands=products.brand.nunique(),primary_categories=snapshots.primary_category.nunique(),retailers=snapshots.retailer_id.nunique(),min=clean.min(),max=clean.max(),mean=clean.mean(),median=clean.median(),q1=q1,q3=q3,iqr=iqr,outliers=int(outlier),complete_price_rate=float(valid.mean()),from_date=str(dates.min()),to_date=str(dates.max()),observed_dates=dates.dt.date.nunique())
    def independent_review(frame):
        ratings=pd.to_numeric(frame.rating,errors='coerce');ratings=ratings[ratings.between(1,5)]
        def rate(column):
            values=frame[column].dropna().str.strip().str.lower();values=values[values.ne('')]
            return float(values.isin(['yes','true','1']).mean()) if len(values) else None
        unknown=frame.product_id_match.fillna('').str.lower().eq('unknown')
        dates=pd.to_datetime(frame.review_date.str.replace(r'^(\d{1,2})\s+tháng\s+(\d{1,2}),?\s+(\d{4})$',r'\3-\2-\1',regex=True),format='mixed',errors='coerce')
        return dict(total=len(frame),products=len(frame.loc[frame.source_product_id.notna() & frame.source_product_id.ne(''),['source','source_product_id']].drop_duplicates()),mean=ratings.mean(),median=ratings.median(),min=ratings.min(),max=ratings.max(),purchased=rate('is_bought'),images=rate('has_images'),unknown=int(unknown.sum()),unknown_rate=float(unknown.mean()),from_date=str(dates.min().date()),to_date=str(dates.max().date()))
    review_result=independent_review(reviews)
    review_by_source={source:independent_review(frame) for source,frame in review_sources.items()}
    evidence=[]
    def check(label,expected,actual):
        passed=close(actual,expected) if isinstance(expected,(int,float)) or hasattr(expected,'item') else actual==expected
        evidence.append(dict(check=label,expected=expected,actual=actual,status='PASS' if passed else 'FAIL'))
        if not passed: raise AssertionError((label,expected,actual))
    eda=load_workbook(HERE/'01_ket_qua_eda.xlsx',read_only=True,data_only=True)
    ns=eda['02_Numeric_Descriptive'];headers=[c.value for c in ns[3]]
    nums={r[0]:dict(zip(headers,r)) for r in ns.iter_rows(min_row=4,values_only=True)}
    for metric in ['min','max','mean','median','q1','q3','iqr']: check('price '+metric,price_result[metric],nums['current_price_vnd'][metric])
    check('price outlier candidates',int(outlier),nums['current_price_vnd']['outlier_candidate_n'])
    check('price valid n',int(valid.sum()),nums['current_price_vnd']['valid_n'])
    for metric in ['min','max','mean','median']: check('review rating '+metric,review_result[metric],nums['rating'][metric])
    for variable,row in nums.items(): check(variable+' partition',row['total_n'],row['valid_n']+row['missing_n']+row['invalid_n'])
    categorical=list(eda['03_Categorical_Descriptive'].iter_rows(min_row=4,values_only=True))
    cat=pd.DataFrame(categorical,columns=[c.value for c in eda['03_Categorical_Descriptive'][3]])
    for (src,var),group in cat.groupby(['source','variable']):
        check(src+'/'+var+' count reconciliation',int(group.denominator_n.iloc[0]),int(group['count'].sum()))
        check(src+'/'+var+' percentages',100.,float(group.percentage.sum()))
    support=eda['06_Review_Descriptive'];stored=dict(zip([c.value for c in support[3]],next(support.iter_rows(min_row=4,max_row=4,values_only=True))))
    review_column={'unknown':'source_flag_unknown_n','unknown_rate':'source_flag_unknown_rate'}
    for metric in ['total','products','mean','median','purchased','images','unknown','unknown_rate']: check('review '+metric,review_result[metric],stored[review_column.get(metric,metric)])
    from stage6_data import read_reviews
    _,prepared=read_reviews()
    for source,expected in review_by_source.items():
        source_row=next(dict(zip([c.value for c in support[3]],r)) for r in support.iter_rows(min_row=4,values_only=True) if r[0]==source)
        for metric in ['total','products','mean','median','purchased','images','unknown','unknown_rate']: check(source+' review '+metric,expected[metric],source_row[review_column.get(metric,metric)])
        parsed=prepared.loc[prepared.source.eq(source),'review_date_parsed']
        check(source+' review date parsing',0,int(parsed.isna().sum()))
        check(source+' earliest review date',expected['from_date'],str(parsed.min().date()))
        check(source+' latest review date',expected['to_date'],str(parsed.max().date()))
    check('source-aware product keys',review_result['products'],prepared.reviewed_product_key.nunique())
    quality_sheet=eda['07_Review_Data_Quality']
    quality_headers=[c.value for c in quality_sheet[3]]
    quality_rows={r[0]:dict(zip(quality_headers,r)) for r in quality_sheet.iter_rows(min_row=4,values_only=True)}
    review_quality_by_source={}
    for source,raw in review_sources.items():
        def blank(series): return series.isna() | series.fillna('').str.strip().eq('')
        def flag_counts(column):
            known=~blank(raw[column]); true=raw[column].fillna('').str.strip().str.lower().isin(['yes','true','1'])
            return int(known.sum()),int(true.sum())
        values=pd.to_numeric(raw.rating,errors='coerce');quality_valid=values.between(1,5)
        dates=pd.to_datetime(raw.review_date.str.replace(r'^(\d{1,2})\s+tháng\s+(\d{1,2}),?\s+(\d{4})$',r'\3-\2-\1',regex=True),format='mixed',errors='coerce')
        status={str(k):int(v) for k,v in raw.product_id_match.where(~blank(raw.product_id_match),'MISSING').value_counts().items()}
        bought_known,bought_true=flag_counts('is_bought');images_known,images_true=flag_counts('has_images')
        expected=dict(source=source,total_reviews=len(raw),unique_source_products=review_by_source[source]['products'],valid_rating_n=int(quality_valid.sum()),missing_rating_n=int(blank(raw.rating).sum()),invalid_rating_n=int((~blank(raw.rating)&~quality_valid).sum()),min_rating=values[quality_valid].min(),max_rating=values[quality_valid].max(),mean_rating=values[quality_valid].mean(),median_rating=values[quality_valid].median(),missing_review_text_n=int(blank(raw.review_text).sum()),missing_review_date_n=int(blank(raw.review_date).sum()),invalid_review_date_n=int((~blank(raw.review_date)&dates.isna()).sum()),duplicate_review_n=int(raw.duplicated(['source','source_product_id','review_id']).sum()),is_bought_known_n=bought_known,is_bought_true_n=bought_true,has_images_known_n=images_known,has_images_true_n=images_true,product_id_match_counts=json.dumps(status,ensure_ascii=False,sort_keys=True),sample_share=len(raw)/len(reviews))
        review_quality_by_source[source]=expected
        for column,value in expected.items(): check(source+' quality '+column,value,quality_rows[source][column])
        check(source+' rating quality partition',len(raw),expected['valid_rating_n']+expected['missing_rating_n']+expected['invalid_rating_n'])
        categorical_status=cat.loc[cat.source.eq(source+'_reviews') & cat.variable.eq('product_id_match')]
        check(source+' raw source flag categorical counts',status,dict(zip(categorical_status.value,categorical_status['count'].astype(int))))
    check('source counts reconcile',len(reviews),sum(r['total_reviews'] for r in review_quality_by_source.values()))
    check('source sample shares reconcile',1.,sum(r['sample_share'] for r in review_quality_by_source.values()))
    main_sheets=['01_Overview','02_Numeric_Descriptive','03_Categorical_Descriptive','04_Data_Coverage_Quality','05_Grouped_Price_Statistics','06_Review_Descriptive','07_Review_Data_Quality','08_Discount_Audit','09_Schema_Inventory']
    check('main EDA sheet order',main_sheets,eda.sheetnames[:9])
    check('legacy summary names removed',False,any(n in eda.sheetnames for n in ['Tong_quan','Thong_ke_mo_ta','Chat_luong_du_lieu','Ty_le_trong','Ngoai_le_gia']))
    check('accepted price numeric values preserved',[r for r in prior_numeric if r['variable']!='rating'],[r for r in numeric_values(ns) if r['variable']!='rating'])
    check('unique collection quality evidence retained',prior_quality,list(eda['11_Collection_Status_Quality'].iter_rows(min_row=4,values_only=True)))
    check('outlier detail evidence retained',prior_outliers,list(eda['12_Price_Outlier_Candidates'].iter_rows(min_row=4,values_only=True)))
    check('extra listed/discount IQR evidence retained',[list(r) for r in prior_extra_iqr],[list(r) for r in eda['10_Extra_IQR_Diagnostics'].iter_rows(min_row=4,values_only=True)])
    if prior_missing:
        files={'products':'products*.csv','price_snapshots':'price_snapshots*.csv','product_matching':'product_matching.csv','collection_logs':'collection_logs.csv','retailers':'retailers.csv','product_categories':'product_categories (*).csv'}
        inventory={(r[0],r[2]):r for r in eda['09_Schema_Inventory'].iter_rows(min_row=4,values_only=True)}
        for src,col,total,count,rate in prior_missing:
            row=inventory[(next(DATA.glob(files[src])).name,col)]
            assert total==row[1] and count==row[4] and close(rate*100,row[5]),(src,col)
        check('legacy field missingness retained in schema inventory',True,True)
    kpi=load_workbook(HERE/'02_kpi_dictionary.xlsx',read_only=True,data_only=True)
    kd={r[0]:r for r in kpi['KPI_dictionary'].iter_rows(min_row=5,values_only=True)}
    check('all 10 core price KPI unchanged',prior_core_kpi,list(kpi['KPI_dictionary'].iter_rows(min_row=5,values_only=True)))
    check('core price KPI count',10,len(kd))
    verified=matching.match_status.fillna('').str.lower().eq('matched') & matching.verified_by_human.astype(str).str.lower().eq('true')
    if not verified.any():
        check('gap unavailable status','NOT CALCULABLE',kd['KPI12'][8]);check('gap is blank',None,kd['KPI12'][7])
    check('matching unavailable status','NOT CALCULABLE',kd['KPI06'][8]);check('matching is blank',None,kd['KPI06'][7])
    listed=pd.to_numeric(snapshots.listed_price_vnd,errors='coerce').replace([math.inf,-math.inf],math.nan)
    pairs=valid & listed.gt(0) & listed.ge(p)
    discount=((listed[pairs]-p[pairs])/listed[pairs]).mean()
    check('recalculated discount',float(discount),kd['KPI09'][7])
    charts=load_workbook(HERE/'02_ket_qua_bieu_do.xlsx')
    check('price chart count',5,len(charts['Bieu_do_phan_tich']._charts))
    check('review chart count',2,len(charts['Review_Charts']._charts))
    testing=load_workbook(HERE/'04_kiem_thu_dashboard.xlsx',read_only=True,data_only=True)
    price_tests=list(testing['Test_case'].iter_rows(min_row=2,values_only=True))
    check('price engine tests all pass',True,all(r[6]=='PASS' for r in price_tests))
    review_tests=list(testing['Review_Test_Cases'].iter_rows(min_row=2,values_only=True))
    check('review engine tests all pass',True,bool(review_tests) and all(r[6]=='PASS' for r in review_tests))
    text=(HERE/'03_insight.md').read_text(encoding='utf-8')
    check('insight count',8,len(re.findall(r'^## I\d+',text,re.M)))
    for label in ['Phát hiện:','Số liệu:','Ý nghĩa:','Hành động:','Bằng chứng:']: check('insight '+label,8,text.count(label))
    page=(HERE/'02_dashboard.html').read_text(encoding='utf-8')
    check('no product filter',False,'id="productSearch"' in page)
    check('review tab present',True,'id="reviewsTab"' in page)
    check('review chart containers',True,'id="ratingChart"' in page and 'id="reviewCategoryChart"' in page)
    # Compile every embedded JS script. Startup UI is not inferred from compilation.
    node=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    scripts=re.findall(r'<script(?![^>]*application/json)[^>]*>(.*?)</script>',page,re.S)
    proc=subprocess.run([str(node),'-e',"const fs=require('fs');JSON.parse(fs.readFileSync(0,'utf8')).forEach(s=>new Function(s));console.log('JavaScript syntax PASS')"],input=json.dumps(scripts),text=True,capture_output=True)
    check('dashboard JS compilation',0,proc.returncode)
    check('raw CSV hash unchanged during full rerun',before,hashes())
    for book in [eda,kpi,charts,testing]:book.close()
    matching_result=dict(candidates=len(matching),verified=int(verified.sum()),cross_source_gap_status=kd['KPI12'][8],matching_rate_status=kd['KPI06'][8])
    if previous:
        check('previous accepted price results retained',previous['price'],price_result)
        if 'review_by_source' in previous:
            check('previous accepted review results retained',previous['review'],review_result)
            check('previous accepted per-source review results retained',previous['review_by_source'],review_by_source)
        else:
            check('previous accepted Guardian review results retained',previous['review'],review_by_source['guardian'])
        check('previous matching results retained',previous['matching'],matching_result)
        check('raw CSV unchanged since previous validation',previous['raw_sha256'],before)
    result=dict(runs=runs,notebook_code_cells_executed=notebooks,notebook_execution_method='Sequential Python AST runner; Jupyter kernel not installed',price=price_result,matching=matching_result,review=review_result,review_by_source=review_by_source,anomalous_discount_n=int((valid & listed.gt(0) & listed.lt(p)).sum()),price_engine_tests=len(price_tests),review_engine_tests=len({r[0] for r in review_tests}),dashboard_ui='NOT_RUN',ui_reason='Engine validation passed; dashboard regenerated, browser QA must be rerun for this build.',evidence=evidence,raw_sha256=before)
    result['review_quality_by_source']=review_quality_by_source
    def scalar(value):
        return value.item() if hasattr(value,'item') else str(value)
    (HERE/'05_validation_results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=scalar),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in ['evidence','raw_sha256']},ensure_ascii=False,indent=2,default=scalar))
    print(f'{len(evidence)} independent validation checks PASS')

if __name__=='__main__':
    if '--startup-only' in sys.argv: check_startup()
    else: main()
