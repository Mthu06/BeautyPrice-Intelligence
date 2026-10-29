"""EDA BEAUTYPRICE: thống kê mô tả, chất lượng dữ liệu và ngoại lệ."""
from pathlib import Path
import math
import sys
import pandas as pd
from stage6_data import (read_reviews, review_metrics, review_source_quality, numeric_description, categorical_description,
                        missing, numeric, unique_nonblank, truth, verified_rows, cross_source_gaps, REVIEW_SCOPE)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
DATA = HERE.parent / "datar"
OUTPUT_FILE = HERE / "01_ket_qua_eda.xlsx"

def source(pattern):
    files = sorted(DATA.glob(pattern))
    if len(files) != 1:
        raise FileNotFoundError(f"Cần đúng một file khớp {pattern}; tìm thấy {len(files)}")
    return files[0]

def read(pattern):
    return pd.read_csv(source(pattern), low_memory=False, encoding="utf-8-sig")

def num(frame, col):
    return pd.to_numeric(frame[col], errors="coerce").replace([float("inf"), -float("inf")], float("nan"))

snapshots = read("price_snapshots*.csv")
products = read("products*.csv")
matching = read("product_matching.csv")
logs = read("collection_logs.csv")
retailers = read("retailers.csv")
categories = read("product_categories (*).csv")

price = num(snapshots, "current_price_vnd")
listed = num(snapshots, "listed_price_vnd")
discount = num(snapshots, "discount_percent")
dates = pd.to_datetime(snapshots["collected_at"], errors="coerce", utc=True).dt.tz_convert("Asia/Ho_Chi_Minh")
valid_price = price.gt(0)
valid_listed = listed.gt(0)
valid_discount = discount.between(0, 100, inclusive="both")
start_date, end_date = dates.min(), dates.max()
period = f"{start_date:%Y-%m-%d} đến {end_date:%Y-%m-%d} (UTC+7)"

def describe(frame, col, unit, rule, valid):
    s = num(frame, col)
    clean = s[valid]
    q1, q3 = clean.quantile(.25), clean.quantile(.75)
    iqr = q3 - q1
    low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    return [col, unit, rule, len(frame), int(valid.sum()), int((~valid).sum()), float((~valid).mean()),
            float(clean.min()), float(clean.max()), float(clean.mean()), float(clean.median()),
            float(clean.std()), float(q1), float(q3), float(iqr), float(low), float(high),
            int(((clean < low) | (clean > high)).sum())]

desc_headers = ["Biến", "Đơn vị", "Điều kiện hợp lệ", "Tổng quan sát", "Hợp lệ", "Thiếu/không hợp lệ",
                "Tỷ lệ thiếu/không hợp lệ", "Giá thấp nhất", "Giá cao nhất", "Trung bình", "Trung vị",
                "Độ lệch chuẩn", "Q1", "Q3", "IQR", "Ngưỡng dưới IQR", "Ngưỡng trên IQR", "Ứng viên ngoại lệ"]
descriptive_rows = [
    describe(snapshots, "current_price_vnd", "VND", "Giá trị số hữu hạn > 0", valid_price),
    describe(snapshots, "listed_price_vnd", "VND", "Giá trị số hữu hạn > 0", valid_listed),
    describe(snapshots, "discount_percent", "%", "Giá trị số trong [0,100]", valid_discount),
]

overview_rows = [
    ["Tập dữ liệu chính", source("price_snapshots*.csv").name],
    ["Đơn vị quan sát", "Một dòng snapshot giá"],
    ["Phạm vi thời gian", period],
    ["Tổng snapshot", len(snapshots)],
    ["Sản phẩm duy nhất trong products", products["product_id"].dropna().nunique()],
    ["Thương hiệu có giá trị", products["brand"].dropna().astype(str).str.strip().replace("", pd.NA).nunique()],
    ["Danh mục trong bảng liên kết", categories["category"].dropna().nunique()],
    ["Nhà bán lẻ", retailers["retailer_id"].dropna().nunique()],
    ["Giá hiện tại hợp lệ", int(valid_price.sum())],
    ["Ngày thu thập thiếu/sai", int(dates.isna().sum())],
    ["Snapshot_id trùng", int(snapshots["snapshot_id"].duplicated().sum())],
]

missing_rows = []
for name, frame in [("products", products), ("price_snapshots", snapshots), ("product_matching", matching), ("collection_logs", logs), ("retailers", retailers), ("product_categories", categories)]:
    for col in frame.columns:
        miss = frame[col].isna() | frame[col].astype(str).str.strip().eq("")
        missing_rows.append([name, col, len(frame), int(miss.sum()), float(miss.mean())])

quality_rows = []
for col in ["price_quality_status", "discount_check_status", "discount_resolution_status", "product_link_status", "retailer_link_status"]:
    for status, count in snapshots[col].fillna("MISSING").value_counts().items():
        quality_rows.append(["price_snapshots", col, status, int(count), count / len(snapshots)])
for status, count in logs["status"].fillna("MISSING").value_counts().items():
    quality_rows.append(["collection_logs", "status", status, int(count), count / len(logs)])

group_headers = ["Nhóm", "Giá trị", "Snapshot", "Giá hợp lệ", "Thiếu/sai", "Tỷ lệ thiếu/sai", "Min (VND)", "Max (VND)", "Trung bình", "Trung vị", "Q1", "Q3", "IQR", "Ứng viên ngoại lệ"]
grouped_rows = []
for col in ["retailer_name", "primary_category", "brand"]:
    for label, g in snapshots.groupby(col, dropna=False):
        s = num(g, "current_price_vnd"); valid = s.gt(0); clean = s[valid]
        q1, q3 = clean.quantile(.25), clean.quantile(.75); iqr = q3 - q1
        grouped_rows.append([col, "Chưa phân loại/thiếu" if pd.isna(label) else str(label), len(g), int(valid.sum()), int((~valid).sum()), float((~valid).mean()), float(clean.min()), float(clean.max()), float(clean.mean()), float(clean.median()), float(q1), float(q3), float(iqr), int(((clean < q1-1.5*iqr) | (clean > q3+1.5*iqr)).sum())])

q1, q3 = price[valid_price].quantile(.25), price[valid_price].quantile(.75)
iqr_price = q3 - q1
outlier_flag = valid_price & ((price < q1-1.5*iqr_price) | (price > q3+1.5*iqr_price))
outlier_rows = snapshots.loc[outlier_flag, ["snapshot_id", "product_id", "retailer_name", "primary_category", "brand", "collected_at", "current_price_vnd", "listed_price_vnd", "source_file"]].values.tolist()

notes_rows = [
    ["Điều kiện lọc giá", "Giá phải là số hữu hạn > 0. Giá trống, sai kiểu, bằng 0 hoặc âm không vào thống kê giá nhưng vẫn nằm trong mẫu số chất lượng."],
    ["Ngoại lệ", "Đánh dấu ngoài [Q1-1,5×IQR; Q3+1,5×IQR]. Không tự xóa; vẫn giữ trong thống kê nếu giá hợp lệ."],
    ["Giải thích", "Giá cao/thấp có thể do thương hiệu, dung tích, quy cách hoặc phân khúc. Cần kiểm tra nguồn trước khi kết luận lỗi."],
    ["Cách xử lý", "Giữ nguyên CSV nguồn, không điền giá thiếu, không winsorize; xuất danh sách ứng viên để rà soát."],
    ["Giới hạn", "Dữ liệu chỉ phản ánh các retailer và đợt thu thập nêu trên, không đại diện toàn thị trường."],
]

def add_table(ws, title, headers, rows, widths=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title; ws["A1"].font = Font(name="Arial", size=15, bold=True, color="17365D")
    for j, h in enumerate(headers, 1):
        c = ws.cell(3, j, h); c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(wrap_text=True)
    for i, row in enumerate(rows, 4):
        for j, value in enumerate(row, 1):
            if isinstance(value, float) and not math.isfinite(value): value = None
            ws.cell(i, j, value)
    ws.freeze_panes = "A4"; ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, len(rows)+3)}"
    for j in range(1, len(headers)+1): ws.column_dimensions[get_column_letter(j)].width = (widths[j-1] if widths else 18)
    for row in ws.iter_rows(min_row=4):
        for c in row:
            if isinstance(c.value, float): c.number_format = "0.00%" if "Tỷ lệ" in str(headers[c.column-1]) else "#,##0.00"

wb = Workbook(); wb.remove(wb.active)
tables = [
    ("01_Overview", "EDA - Phạm vi và tập dữ liệu", ["Chỉ tiêu", "Giá trị"], overview_rows, [34, 100]),
    ("10_Extra_IQR_Diagnostics", "Supporting IQR diagnostics; main statistics in 02_Numeric_Descriptive", ["variable", "unit", "validity_rule", "lower_bound", "upper_bound", "outlier_candidate_n"], [[r[0],r[1],r[2],r[15],r[16],r[17]] for r in descriptive_rows[1:]], [30,14,55,22,22,24]),
    ("11_Collection_Status_Quality", "Collection and source validation statuses", ["Nguồn", "Trường trạng thái", "Giá trị", "Số lượng", "Tỷ lệ"], quality_rows, [24, 30, 36, 14, 14]),
    ("05_Grouped_Price_Statistics", "EDA - Thống kê giá theo retailer, danh mục và thương hiệu", group_headers, grouped_rows, None),
    ("12_Price_Outlier_Candidates", "EDA - Ứng viên ngoại lệ giá cần rà soát", ["snapshot_id", "product_id", "retailer", "danh mục", "thương hiệu", "collected_at", "current_price_vnd", "listed_price_vnd", "source_file"], outlier_rows, None),
    ("13_Method_Notes", "EDA - Quy tắc và nhận xét", ["Nội dung", "Giải thích"], notes_rows, [32, 120]),
]
for name, title, headers, rows, widths in tables:
    add_table(wb.create_sheet(name), title, headers, rows, widths)

# Extend the original EDA with explicit numeric/categorical/coverage outputs.
review_paths, reviews = read_reviews()
review_summary = review_metrics(reviews)
review_summary_rows = [{'source':'all',**review_summary,'sample_share':1.0}] + [{'source':source,**review_metrics(frame),'sample_share':len(frame)/len(reviews)} for source,frame in reviews.groupby('source')]
review_source_quality_frame = review_source_quality(reviews)
pair_positive = valid_price & valid_listed
anomalous_discount = pair_positive & listed.lt(price)
valid_discount_pair = pair_positive & ~anomalous_discount
snapshots['discount_percent_recalculated'] = ((listed-price)/listed*100).where(pair_positive)
snapshots['anomalous_discount'] = anomalous_discount
numeric_records = []
for column, validity, unit, rule in [
    ('current_price_vnd', valid_price, 'VND', 'Finite current price > 0; IQR candidates retained'),
    ('listed_price_vnd', valid_listed, 'VND', 'Finite listed price > 0'),
    ('discount_percent', valid_discount, '%', 'Resolved input discount in [0,100]; original website value is discount_percent_source'),
    ('discount_percent_calculated', numeric(snapshots['discount_percent_calculated']).between(0,100), '%', 'Source calculated discount in [0,100]'),
    ('discount_percent_recalculated', valid_discount_pair, '%', 'P>0; L>0; L>=P; exclude anomalous negative discounts'),
]:
    numeric_records.append(numeric_description(column, source('price_snapshots*.csv').name, snapshots[column], validity, unit, rule, column=='current_price_vnd'))
numeric_records.append(numeric_description('rating', 'hasaki_reviews.csv + guardian_reviews.csv', reviews['rating'], numeric(reviews['rating']).between(1,5), 'stars', 'All source review rows; range check 1–5; half stars retained'))
numeric_frame = pd.DataFrame(numeric_records)
categorical_records = []
categorical_records += categorical_description(snapshots, ['brand','primary_category','retailer_name'], 'price_snapshots')
categorical_records += categorical_description(products, ['brand','primary_category'], 'products')
categorical_records += categorical_description(matching, ['match_status','verified_by_human'], 'product_matching')
for source_name,frame in reviews.groupby('source'):
    categorical_records += categorical_description(frame, ['content_quality','is_bought','has_images','product_id_match','category'], source_name+'_reviews')
categorical_frame = pd.DataFrame(categorical_records)
verified = verified_rows(matching)
gaps = cross_source_gaps(snapshots, matching)
coverage = []
def coverage_item(scope, metric, value, definition):
    coverage.append([scope, metric, value, definition])
for metric, value, definition in [
    ('total_price_snapshots',len(snapshots),'All input snapshot rows'),
    ('unique_product_id',unique_nonblank(snapshots['product_id']),'Unique snapshot products'),
    ('master_unique_products',unique_nonblank(products['product_id']),'Products master; distinct from snapshot products'),
    ('unique_brands',unique_nonblank(snapshots['brand']),'Nonblank snapshot brand strings; no alias normalization'),
    ('unique_primary_categories',unique_nonblank(snapshots['primary_category']),'Nonblank primary_category; missing shown separately'),
    ('unique_linked_categories',unique_nonblank(categories['category']),'Distinct labels in category relationship table'),
    ('unique_retailers',unique_nonblank(snapshots['retailer_id']),'Distinct retailers observed'),
    ('earliest_collected_at',str(dates.min()),'UTC+7'),('latest_collected_at',str(dates.max()),'UTC+7'),
    ('observed_dates',dates.dt.date.nunique(),'Distinct observed local dates'),
    ('duplicate_snapshot_id',int(snapshots['snapshot_id'].duplicated().sum()),'Excess duplicate IDs'),
    ('duplicate_full_rows',int(snapshots.drop(columns=['discount_percent_recalculated','anomalous_discount']).duplicated().sum()),'Excess identical rows'),
    ('missing_current_price',int(missing(snapshots['current_price_vnd']).sum()),'Blank/NA only'),
    ('invalid_current_price',int((~missing(snapshots['current_price_vnd']) & ~valid_price).sum()),'Nonblank nonnumeric/nonfinite/<=0'),
    ('missing_listed_price',int(missing(snapshots['listed_price_vnd']).sum()),'Blank/NA only'),
    ('missing_discount',int(missing(snapshots['discount_percent']).sum()),'Blank/NA only'),
    ('anomalous_discount',int(anomalous_discount.sum()),'P>0 and L>0 but L<P; excluded from valid discount'),
    ('current_price_outlier_candidates',int(outlier_flag.sum()),'IQR candidates retained'),
    ('complete_price_rate_pct',float(valid_price.mean()*100),'Valid price / all snapshots * 100'),
    ('overall_observed_price_range_vnd',float(price[valid_price].max()-price[valid_price].min()),'Descriptive catalog range only; not cross-source gap'),
]: coverage_item('PRICE',metric,value,definition)
for metric, value in [('matching_candidates',len(matching)),('verified_by_human_true',int(truth(matching['verified_by_human']).sum())),('verified_by_human_false',int(matching['verified_by_human'].astype(str).str.lower().eq('false').sum())),('matched_and_verified_rows',len(verified)),('verified_comparable_product_dates',len(gaps))]:
    coverage_item('MATCHING',metric,value,'Observed matching file; not a formal matching denominator')
for metric, value in review_summary.items(): coverage_item('REVIEW',metric,value,REVIEW_SCOPE)
for row in categorical_records:
    if row['source'] in ['product_matching','guardian_reviews','hasaki_reviews']:
        coverage_item('MATCHING' if row['source']=='product_matching' else 'REVIEW',row['variable']+'='+row['value'],row['count'],f"Denominator={row['denominator_n']}; percentage={row['percentage']:.4f}%; missing retained")
for metric, value, rule in [
    ('earliest_review_date',str(reviews['review_date_parsed'].min().date()),'Source review date; not collection date'),
    ('latest_review_date',str(reviews['review_date_parsed'].max().date()),'Source review date; not collection date'),
    ('missing_rating',int(missing(reviews['rating']).sum()),'Blank/NA only'),
    ('missing_review_text',int(missing(reviews['review_text']).sum()),'Blank/NA only'),
    ('missing_review_date',int(missing(reviews['review_date']).sum()),'Blank/NA only'),
    ('invalid_review_date',int((~missing(reviews['review_date']) & reviews['review_date_parsed'].isna()).sum()),'Nonblank parse failures'),
    ('duplicate_reviews',int(reviews.duplicated(['source','source_product_id','review_id']).sum()),'Excess duplicates using source+product+review ID'),
    ('duplicate_flag_true',int(truth(reviews['is_duplicate']).sum()),'Raw duplicate flag; no rows removed'),
]: coverage_item('REVIEW',metric,value,rule)
coverage_frame = pd.DataFrame(coverage,columns=['scope','metric','value','definition'])
review_quality = categorical_frame.loc[categorical_frame.source.isin(['guardian_reviews','hasaki_reviews'])].copy()
schema_rows = []
discount_audit=snapshots[['snapshot_id','product_id','retailer_name','collected_at','current_price_vnd','listed_price_vnd','discount_percent_source','discount_percent','discount_percent_calculated','discount_percent_recalculated','anomalous_discount']].copy()
discount_audit['valid_discount']=valid_discount_pair
discount_audit['source_minus_recalculated_pp']=numeric(snapshots['discount_percent_source'])-snapshots['discount_percent_recalculated']
for path in sorted(DATA.rglob('*.csv')):
    frame = pd.read_csv(path, low_memory=False, encoding='utf-8-sig')
    for col in frame:
        schema_rows.append([str(path.relative_to(DATA)),len(frame),col,str(frame[col].dtype),int(missing(frame[col]).sum()),float(missing(frame[col]).mean()*100)])
for name,title,frame,widths in [
    ('02_Numeric_Descriptive','Numeric descriptive statistics',numeric_frame,None),
    ('03_Categorical_Descriptive','Categorical statistics; percentage denominator = source rows',categorical_frame,[24,26,38,14,16,16,18]),
    ('04_Data_Coverage_Quality','Data coverage and quality',coverage_frame,[18,40,38,90]),
    ('06_Review_Descriptive',REVIEW_SCOPE,pd.DataFrame(review_summary_rows).rename(columns={'unknown':'source_flag_unknown_n','unknown_rate':'source_flag_unknown_rate'}),None),
    ('07_Review_Data_Quality','Source-level review quality; all source rows retained',review_source_quality_frame,None),
    ('09_Schema_Inventory','Schema and missingness of all actual CSV inputs',pd.DataFrame(schema_rows,columns=['file','rows','column','dtype','missing_n','missing_rate_pct']),[65,14,38,18,18,22]),
    ('08_Discount_Audit','Derived discount audit; raw inputs preserved',discount_audit,None),
]:
    add_table(wb.create_sheet(name),title,list(frame.columns),frame.astype(object).where(frame.notna(),None).values.tolist(),widths)
    sheet=wb[name]
    sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(frame.columns))
    sheet.row_dimensions[3].height=44
    for row in sheet.iter_rows(min_row=4):
        sheet.row_dimensions[row[0].row].height=42 if name!='08_Discount_Audit' else 30
        for cell in row:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,float): cell.number_format='#,##0.00'
    if name=='02_Numeric_Descriptive':
        for column in ['total_n','valid_n','missing_n','invalid_n','outlier_candidate_n']:
            j=list(frame.columns).index(column)+1
            for i in range(4,len(frame)+4): sheet.cell(i,j).number_format='#,##0'
        sheet.column_dimensions['A'].width=32
        sheet.column_dimensions['B'].width=28
        sheet.column_dimensions['U'].width=65
        sheet.row_dimensions[3].height=56
    if name=='06_Review_Descriptive':
        for column in ['purchased','images','source_flag_unknown_rate','sample_share']:
            j=list(frame.columns).index(column)+1
            for i in range(4,len(frame)+4): sheet.cell(i,j).number_format='0.00%'
        sheet['A2']='Source product-match flag: this source field is NOT a verified link to price products.'
        sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(frame.columns))
        sheet.row_dimensions[2].height=30
        sheet['A2'].alignment=Alignment(wrap_text=True,vertical='top')
        for col in ['H','I']: sheet.column_dimensions[col].width=28
    if name=='07_Review_Data_Quality':
        sheet['A2']='Observed rating distribution in the collected sample. Duplicate = excess rows by source + source_product_id + review_id. Source flag categorical counts also retained in 03_Categorical_Descriptive; no verified price link.'
        sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(frame.columns))
        sheet.row_dimensions[2].height=34
        sheet['A2'].alignment=Alignment(wrap_text=True,vertical='top')
        sheet.row_dimensions[3].height=65
        for j,column in enumerate(frame.columns,1):
            sheet.column_dimensions[get_column_letter(j)].width=24 if column!='product_id_match_counts' else 34
            for i in range(4,len(frame)+4):
                sheet.cell(i,j).number_format='0.00%' if column=='sample_share' else '0.0000' if column in ['mean_rating','median_rating','min_rating','max_rating'] else '#,##0' if column.endswith('_n') or column in ['total_reviews','unique_source_products'] else 'General'
notes=wb['13_Method_Notes']
notes.append(['Review scope',REVIEW_SCOPE])
notes.append(['Review mapping','All review metrics independent of price; no verified product join. Source identifiers may distinguish variants/combos.'])
notes.append(['Source product-match flag','This source field is NOT a verified link to price products. Hasaki source match and Guardian unknown are reported literally in source-specific categorical counts, not a verified matching rate.'])
notes.append(['Review sample imbalance','Pooled review metrics weight each review equally; source shares are in 06 and 07. Compare per-source metrics, not pooled average, when describing the two collected samples; do not infer retailer customer satisfaction.'])
notes.append(['Rating scale','Hasaki and Guardian: validate 1–5, retain half stars; no rescaling. Dates parsed with source-specific formats. Products counted by source + source_product_id; original category labels retained.'])
notes.append(['discount_percent','Resolved discount stored in input CSV: source_value_validated keeps the website value; replaced_by_calculated_value_source_preserved replaces a mismatch; filled_from_valid_price_pair fills a missing source discount. Not always the original website value.'])
notes.append(['discount_percent_source','Original website/source discount retained in input CSV; compare with recalculation to audit source discrepancies.'])
notes.append(['discount_percent_calculated','Pre-existing calculated field in input CSV; actual values match (L-P)/L*100 rounded to 2 decimals. Upstream processing code is not included, so the exact upstream step is not asserted.'])
notes.append(['discount_percent_recalculated','Stage 6 independent dataframe calculation (L-P)/L*100, no rounding before aggregation. Valid average uses P>0, L>0, L>=P. Negative values may be kept in audit with anomalous_discount=True, excluded from valid numeric statistics/KPI.'])
notes.append(['Workbook cleanup','Legacy 3-variable summary is superseded by 02_Numeric_Descriptive. Unique listed-price/discount IQR diagnostics are retained in 10_Extra_IQR_Diagnostics; field missingness moved to 09_Schema_Inventory.'])
for i in range(4,notes.max_row+1):
    notes.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top')
    notes.row_dimensions[i].height=56
wb['01_Overview'].append(['Cách đọc','Thống kê mô tả chính: 02_Numeric_Descriptive (6 biến). Sheet 10–13 chỉ chứa chẩn đoán và ghi chú bổ trợ.'])
wb['01_Overview'].cell(wb['01_Overview'].max_row,2).alignment=Alignment(wrap_text=True)
wb['01_Overview'].row_dimensions[wb['01_Overview'].max_row].height=35
# Main nine worksheets first; unique supporting evidence retained with explicit names.
wb._sheets=sorted(wb._sheets,key=lambda sheet:sheet.title)
wb.save(OUTPUT_FILE)
print(f"Created: {OUTPUT_FILE}")
print(f"Giá hợp lệ: {valid_price.sum():,}; min={price[valid_price].min():,.0f}; max={price[valid_price].max():,.0f}; ngoại lệ IQR={outlier_flag.sum():,}")
